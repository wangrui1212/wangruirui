#!/usr/bin/env python
"""Update the confirmed workbook with continuous coverage from the analysis manifest."""

from __future__ import annotations

import json
from copy import copy
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.styles import PatternFill


PROJECT_ROOT = Path(__file__).resolve().parents[1]
WORKBOOK = (
    PROJECT_ROOT
    / "井中流动分析"
    / "PSD分析_20260710"
    / "中国石油大学PSD分析法_实验记录与数据覆盖_已确认.xlsx"
)
MANIFEST = (
    PROJECT_ROOT
    / "井中流动分析"
    / "PSD分析_20260710"
    / "三设备分析结果"
    / "PSD_20260710_三设备分析清单.json"
)
OUTPUT_WORKBOOK = WORKBOOK.with_name("中国石油大学PSD分析法_实验记录与数据覆盖_连续覆盖复核版.xlsx")


def main() -> int:
    rows = json.loads(MANIFEST.read_text(encoding="utf-8"))
    lookup = {
        (row["experiment_type"], row["experiment_id"], row["device"]): row
        for row in rows
    }
    workbook = load_workbook(WORKBOOK)
    sheet = workbook["数据覆盖检查"]
    headers = {cell.value: cell.column for cell in sheet[1]}
    coverage_col = sheet.max_column + 1
    missing_col = sheet.max_column + 2
    sheet.cell(1, coverage_col, "连续覆盖率(%)")
    sheet.cell(1, missing_col, "缺失时段")
    for column in (coverage_col, missing_col):
        sheet.cell(1, column).fill = PatternFill("solid", fgColor="1F4E78")
        sheet.cell(1, column).font = copy(sheet.cell(1, 1).font)
        sheet.cell(1, column).alignment = copy(sheet.cell(1, 1).alignment)

    partial_fill = PatternFill("solid", fgColor="FCE4D6")
    complete_fill = PatternFill("solid", fgColor="E2F0D9")
    for row_index in range(2, sheet.max_row + 1):
        key = (
            sheet.cell(row_index, headers["实验类型"]).value,
            sheet.cell(row_index, headers["实验编号"]).value,
            sheet.cell(row_index, headers["设备"]).value,
        )
        result = lookup[key]
        coverage = float(result["coverage_percent"])
        missing = result["missing_intervals"]
        status = "完整覆盖" if result["status"] == "ok" else "部分覆盖"
        note = f"按文件头持续时间与实验组起止时间计算，连续覆盖率 {coverage:.2f}%"
        if missing:
            note += f"；缺失 {missing}"
        sheet.cell(row_index, headers["覆盖状态"], status)
        sheet.cell(row_index, headers["说明"], note)
        sheet.cell(row_index, coverage_col, coverage)
        sheet.cell(row_index, missing_col, missing)
        fill = complete_fill if status == "完整覆盖" else partial_fill
        for cell in sheet[row_index]:
            cell.fill = fill

    sheet.column_dimensions[sheet.cell(1, coverage_col).column_letter].width = 16
    sheet.column_dimensions[sheet.cell(1, missing_col).column_letter].width = 28

    qc = workbook["人工核对清单"]
    for row in range(2, qc.max_row + 1):
        if qc.cell(row, 2).value == "原始数据覆盖":
            qc.cell(row, 3, "已复核")
            qc.cell(row, 4, "20个组合完整覆盖；10个组合部分覆盖，缺失区间见数据覆盖检查sheet。")
            break
    workbook.save(OUTPUT_WORKBOOK)
    print(OUTPUT_WORKBOOK)
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
