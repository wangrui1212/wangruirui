#!/usr/bin/env python
"""Batch export full-position phase heatmaps for the well-flow experiment table.

The data reader and display aggregation are reused from tdas_dat_viewer.py.
This module only wraps batch matching, raster rendering, and metadata export.
"""

from __future__ import annotations

import argparse
import csv
import json
import re
import sys
import types
from dataclasses import dataclass, asdict
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Iterable

import numpy as np
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = PROJECT_ROOT / "井中流动分析" / "井中流动分析_两组实验记录.xlsx"
DEFAULT_DATA_ROOT = Path(r"E:\北京试验\ov\2026-07-09")
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "井中流动分析" / "相位热力图"
DEFAULT_VLIM = 0.3
DEFAULT_MAX_TIME_BINS = 4096


def install_gui_stubs() -> None:
    """Let tdas_dat_viewer import in non-GUI batch environments."""
    if "pyqtgraph" not in sys.modules:
        sys.modules["pyqtgraph"] = types.ModuleType("pyqtgraph")
    if "pyqtgraph.opengl" not in sys.modules:
        sys.modules["pyqtgraph.opengl"] = types.ModuleType("pyqtgraph.opengl")
    if "PyQt6" not in sys.modules:
        sys.modules["PyQt6"] = types.ModuleType("PyQt6")
    for name in ("QtCore", "QtGui", "QtWidgets"):
        module_name = f"PyQt6.{name}"
        if module_name not in sys.modules:
            sys.modules[module_name] = types.ModuleType(module_name)
    sys.modules["PyQt6.QtWidgets"].QMainWindow = object


install_gui_stubs()
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from tdas_dat_viewer import (  # noqa: E402
    DatHeader,
    make_display_matrix_max_abs,
    normalize_silixa_matrix,
    read_header,
    read_signal_matrix,
)


@dataclass(frozen=True)
class ExperimentRecord:
    group: str
    date_text: str
    seq: str
    start: datetime
    middle_text: str
    end: datetime
    end_text: str
    operation: str
    flow_lh: float
    unit: str

    @property
    def label(self) -> str:
        flow_text = f"{self.flow_lh:g}{self.unit.replace('/', '')}"
        return f"{self.group}_{self.seq}_{self.start:%Y%m%d_%H%M%S}-{self.end:%H%M%S}_{flow_text}"


@dataclass(frozen=True)
class DatFile:
    path: Path
    timestamp: datetime


@dataclass
class ExportSummary:
    group: str
    seq: str
    start: str
    end: str
    operation: str
    flow_lh: float
    file_count: int
    status: str
    first_file: str
    last_file: str
    protocol: str = ""
    frequency_hz: int | str = ""
    pulse_width_ns: float | str = ""
    point_count: int | str = ""
    position_start_m: float | str = ""
    position_end_m: float | str = ""
    sample_count: int | str = ""
    duration_s: float | str = ""
    image_path: str = ""
    message: str = ""


def as_text(value: object) -> str:
    if value is None:
        return ""
    return str(value).strip()


def parse_date(value: object) -> date:
    text = as_text(value)
    match = re.search(r"\d{4}-\d{2}-\d{2}", text)
    if not match:
        raise ValueError(f"无法解析日期: {text}")
    return date.fromisoformat(match.group(0))


def parse_time_value(value: object, use_last: bool = False) -> time:
    text = as_text(value)
    matches = re.findall(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", text)
    if not matches:
        raise ValueError(f"无法解析时间: {text}")
    hour, minute, second = matches[-1 if use_last else 0]
    return time(int(hour), int(minute), int(second or 0))


def read_records(workbook_path: Path) -> list[ExperimentRecord]:
    workbook = load_workbook(workbook_path, data_only=True)
    records: list[ExperimentRecord] = []
    for sheet_name, group in (("第一组实验", "第一组实验"), ("第二组实验", "第二组实验")):
        if sheet_name not in workbook.sheetnames:
            continue
        sheet = workbook[sheet_name]
        for row in sheet.iter_rows(min_row=4, values_only=True):
            if not row or not row[0]:
                continue
            test_date = parse_date(row[0])
            start_time = parse_time_value(row[2])
            end_time = parse_time_value(row[4], use_last=True)
            start_dt = datetime.combine(test_date, start_time)
            end_dt = datetime.combine(test_date, end_time)
            if end_dt <= start_dt:
                raise ValueError(f"{sheet_name} {row[1]} 结束时间不晚于开始时间")
            records.append(
                ExperimentRecord(
                    group=group,
                    date_text=as_text(row[0]),
                    seq=as_text(row[1]),
                    start=start_dt,
                    middle_text=as_text(row[3]),
                    end=end_dt,
                    end_text=as_text(row[4]),
                    operation=as_text(row[5]),
                    flow_lh=float(row[6]),
                    unit=as_text(row[7]) or "L/h",
                )
            )
    return apply_confirmed_well_flow_times(records)


def apply_confirmed_well_flow_times(records: list[ExperimentRecord]) -> list[ExperimentRecord]:
    """Use user-confirmed rows when Excel time serialization is ambiguous."""
    corrected: list[ExperimentRecord] = []
    for record in records:
        if record.group == "第一组实验" and record.seq == "1" and int(record.flow_lh) == 260:
            corrected.append(
                ExperimentRecord(
                    group=record.group,
                    date_text=record.date_text,
                    seq=record.seq,
                    start=datetime.combine(record.start.date(), time(15, 27)),
                    middle_text="15:38",
                    end=datetime.combine(record.start.date(), time(15, 43)),
                    end_text="15:43",
                    operation=record.operation,
                    flow_lh=record.flow_lh,
                    unit=record.unit,
                )
            )
            continue
        if record.group == "第二组实验" and record.seq == "E" and int(record.flow_lh) == 250:
            corrected.append(
                ExperimentRecord(
                    group=record.group,
                    date_text=record.date_text,
                    seq=record.seq,
                    start=datetime.combine(record.start.date(), time(17, 5)),
                    middle_text="17:10",
                    end=datetime.combine(record.start.date(), time(17, 20)),
                    end_text="17:15-17:20",
                    operation=record.operation,
                    flow_lh=record.flow_lh,
                    unit=record.unit,
                )
            )
            continue
        corrected.append(record)
    return corrected


def parse_file_timestamp(path: Path) -> datetime | None:
    name = path.name
    patterns = [
        (r"(\d{4})-(\d{2})-(\d{2})-(\d{2})-(\d{2})-(\d{2})-out\.dat$", "%Y-%m-%d-%H-%M-%S"),
        (r"(\d{14})\.\d+\.dat$", "%Y%m%d%H%M%S"),
        (r"_(\d{8})_(\d{6})(?:\.|_)", "%Y%m%d%H%M%S"),
    ]
    for pattern, fmt in patterns:
        match = re.search(pattern, name, re.IGNORECASE)
        if not match:
            continue
        if len(match.groups()) == 1:
            stamp = match.group(1)
        elif fmt == "%Y%m%d%H%M%S":
            stamp = match.group(1) + match.group(2)
        else:
            stamp = "-".join(match.groups())
        try:
            return datetime.strptime(stamp, fmt)
        except ValueError:
            continue
    return None


def index_data_files(data_root: Path) -> list[DatFile]:
    patterns = ("*.dat", "*.DAT", "*.tdms", "*.TDMS")
    files: list[DatFile] = []
    seen_paths: set[str] = set()
    for pattern in patterns:
        for path in data_root.rglob(pattern):
            path_key = str(path.resolve()).casefold()
            if path_key in seen_paths:
                continue
            seen_paths.add(path_key)
            timestamp = parse_file_timestamp(path)
            if timestamp is not None:
                files.append(DatFile(path=path, timestamp=timestamp))
    return sorted(files, key=lambda item: item.timestamp)


def files_for_record(files: list[DatFile], record: ExperimentRecord) -> list[DatFile]:
    return [item for item in files if record.start <= item.timestamp <= record.end]


def validate_headers(paths: list[Path]) -> tuple[DatHeader, dict[Path, DatHeader]]:
    headers = {path: read_header(path) for path in paths}
    base = headers[paths[0]]
    total_samples = 0
    for path in paths:
        header = headers[path]
        if header.protocol != base.protocol:
            raise ValueError(f"协议类型不同: {path.name}={header.protocol}, {paths[0].name}={base.protocol}")
        if header.point_count != base.point_count:
            raise ValueError(f"点位数不同: {path.name}={header.point_count}, {paths[0].name}={base.point_count}")
        if abs(header.position_start_m - base.position_start_m) > 1e-6:
            raise ValueError(f"起始位置不同: {path.name}={header.position_start_m}, {paths[0].name}={base.position_start_m}")
        if abs(header.position_end_m - base.position_end_m) > 1e-6:
            raise ValueError(f"结束位置不同: {path.name}={header.position_end_m}, {paths[0].name}={base.position_end_m}")
        if abs(header.point_interval_m - base.point_interval_m) > 1e-6:
            raise ValueError(f"点位间隔不同: {path.name}={header.point_interval_m}, {paths[0].name}={base.point_interval_m}")
        if header.frequency_hz != base.frequency_hz:
            raise ValueError(f"采样率不同: {path.name}={header.frequency_hz}, {paths[0].name}={base.frequency_hz}")
        if header.protocol == "CDAS" and header.channel != base.channel:
            raise ValueError(f"通道号不同: {path.name}=通道{header.channel}, {paths[0].name}=通道{base.channel}")
        total_samples += header.sample_count

    combined = DatHeader(
        protocol=base.protocol,
        system_id=base.system_id,
        year=base.year,
        month=base.month,
        day=base.day,
        hour=base.hour,
        minute=base.minute,
        second=base.second,
        length=base.length,
        global_alarm=base.global_alarm,
        acquisition_status=base.acquisition_status,
        channel=base.channel,
        frequency_hz=base.frequency_hz,
        acquisition_time_s=total_samples / max(1, base.frequency_hz),
        point_interval_m=base.point_interval_m,
        point_count=base.point_count,
        sample_count=total_samples,
        position_start_m=base.position_start_m,
        raw_value_note=base.raw_value_note,
        header_size=base.header_size,
        bytes_per_sample=base.bytes_per_sample,
        timestamp_text=base.timestamp_text,
        extra_info=base.extra_info,
    )
    return combined, headers


def load_display_matrix(paths: list[Path], max_time_bins: int) -> tuple[np.ndarray, DatHeader]:
    header, headers = validate_headers(paths)
    if header.sample_count <= max_time_bins or header.protocol == "Silixa TDMS":
        matrices = []
        for path in paths:
            matrices.append(read_signal_matrix(path, headers[path]))
        matrix = matrices[0] if len(matrices) == 1 else np.concatenate(matrices, axis=1)
        if header.protocol == "Silixa TDMS":
            matrix, _max_abs = normalize_silixa_matrix(matrix)
        display = make_display_matrix_max_abs(matrix, max_time_bins=max_time_bins)
        return display, header

    # Equivalent to make_display_matrix_max_abs(np.concatenate(...), max_time_bins),
    # but streams one file at a time so long intervals do not need multi-GB RAM.
    edges = np.linspace(0, header.sample_count, max_time_bins + 1, dtype=np.int64)
    display = np.zeros((header.point_count, max_time_bins), dtype=np.float32)
    best_abs = np.full((header.point_count, max_time_bins), -np.inf, dtype=np.float32)
    row_indices = np.arange(header.point_count)
    sample_offset = 0
    for path in paths:
        file_header = headers[path]
        matrix = read_signal_matrix(path, file_header)
        file_start = sample_offset
        file_stop = sample_offset + file_header.sample_count
        first_bin = max(0, int(np.searchsorted(edges, file_start, side="right") - 1))
        last_bin = min(max_time_bins - 1, int(np.searchsorted(edges, file_stop - 1, side="right") - 1))
        for bin_index in range(first_bin, last_bin + 1):
            global_start = int(edges[bin_index])
            global_stop = int(edges[bin_index + 1])
            overlap_start = max(global_start, file_start)
            overlap_stop = min(global_stop, file_stop)
            if overlap_stop <= overlap_start:
                continue
            local_start = overlap_start - file_start
            local_stop = overlap_stop - file_start
            block = matrix[:, local_start:local_stop]
            max_indices = np.argmax(np.abs(block), axis=1)
            values = block[row_indices, max_indices]
            abs_values = np.abs(values)
            update_mask = abs_values > best_abs[:, bin_index]
            display[update_mask, bin_index] = values[update_mask]
            best_abs[update_mask, bin_index] = abs_values[update_mask]
        sample_offset = file_stop
    return display, header


def find_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def colorize(values: np.ndarray, vlim: float) -> Image.Image:
    clipped = np.clip(values, -vlim, vlim)
    normalized = (clipped + vlim) / (2.0 * vlim)
    anchors = np.array(
        [
            [33, 102, 172],
            [247, 247, 247],
            [178, 24, 43],
        ],
        dtype=np.float32,
    )
    lower = normalized <= 0.5
    rgb = np.empty(values.shape + (3,), dtype=np.float32)
    left_t = np.clip(normalized / 0.5, 0.0, 1.0)
    right_t = np.clip((normalized - 0.5) / 0.5, 0.0, 1.0)
    rgb[lower] = anchors[0] + (anchors[1] - anchors[0]) * left_t[lower, None]
    rgb[~lower] = anchors[1] + (anchors[2] - anchors[1]) * right_t[~lower, None]
    return Image.fromarray(np.uint8(np.clip(rgb, 0, 255)), mode="RGB")


def draw_centered_text(draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, font: ImageFont.ImageFont, fill: str) -> None:
    bbox = draw.textbbox((0, 0), text, font=font)
    width = bbox[2] - bbox[0]
    draw.text((xy[0] - width // 2, xy[1]), text, font=font, fill=fill)


def render_heatmap(
    display: np.ndarray,
    header: DatHeader,
    record: ExperimentRecord,
    paths: list[Path],
    output_path: Path,
    vlim: float,
) -> None:
    image_matrix = display.T
    heatmap = colorize(image_matrix, vlim=vlim)

    canvas_w, canvas_h = 1800, 1120
    left, top, plot_w, plot_h = 145, 165, 1360, 760
    right = left + plot_w
    bottom = top + plot_h
    canvas = Image.new("RGB", (canvas_w, canvas_h), "white")
    draw = ImageDraw.Draw(canvas)

    title_font = find_font(46)
    label_font = find_font(24)
    small_font = find_font(20)
    tick_font = find_font(18)

    title = f"{record.group} {record.seq} | {record.start:%H:%M:%S}-{record.end:%H:%M:%S} | {record.operation} {record.flow_lh:g}{record.unit}"
    subtitle = f"相位热力图 | 全点位空间 {header.position_start_m:.1f}-{header.position_end_m:.1f} m | 色条 ±{vlim:g} | 文件 {len(paths)} 个"
    draw.text((70, 38), title, font=title_font, fill="#111111")
    draw.text((72, 100), subtitle, font=label_font, fill="#444444")

    heatmap = heatmap.resize((plot_w, plot_h), resample=Image.Resampling.BILINEAR)
    canvas.paste(heatmap, (left, top))
    draw.rectangle((left, top, right, bottom), outline="#111111", width=2)

    x_ticks = np.linspace(header.position_start_m, header.position_end_m, 6)
    for tick in x_ticks:
        x = left + int(round((tick - header.position_start_m) / max(header.distance_span_m, header.point_interval_m) * plot_w))
        draw.line((x, bottom, x, bottom + 8), fill="#111111", width=2)
        draw_centered_text(draw, (x, bottom + 12), f"{tick:.0f}", tick_font, "#111111")

    y_ticks = np.linspace(0.0, header.time_span_s, 6)
    for tick in y_ticks:
        y = top + int(round(tick / max(header.time_span_s, 1e-9) * plot_h))
        draw.line((left - 8, y, left, y), fill="#111111", width=2)
        bbox = draw.textbbox((0, 0), f"{tick:.0f}", font=tick_font)
        draw.text((left - 16 - (bbox[2] - bbox[0]), y - 10), f"{tick:.0f}", font=tick_font, fill="#111111")

    draw_centered_text(draw, ((left + right) // 2, bottom + 52), "Distance (m)", label_font, "#111111")
    draw.text((36, (top + bottom) // 2 - 18), "Time (s)", font=label_font, fill="#111111")

    bar_x, bar_y, bar_w, bar_h = right + 72, top, 42, plot_h
    ramp_values = np.linspace(vlim, -vlim, bar_h, dtype=np.float32).reshape(bar_h, 1)
    bar = colorize(ramp_values, vlim=vlim).resize((bar_w, bar_h), resample=Image.Resampling.BILINEAR)
    canvas.paste(bar, (bar_x, bar_y))
    draw.rectangle((bar_x, bar_y, bar_x + bar_w, bar_y + bar_h), outline="#111111", width=1)
    for value in (vlim, 0.0, -vlim):
        y = bar_y + int(round((vlim - value) / (2.0 * vlim) * bar_h))
        draw.line((bar_x + bar_w, y, bar_x + bar_w + 8, y), fill="#111111", width=2)
        draw.text((bar_x + bar_w + 14, y - 12), f"{value:.1f}", font=tick_font, fill="#111111")
    draw.text((bar_x - 8, bar_y - 36), "Phase", font=label_font, fill="#111111")

    file_text = f"{paths[0].name}  →  {paths[-1].name}"
    draw.text((70, canvas_h - 92), file_text, font=small_font, fill="#333333")
    draw.text((70, canvas_h - 54), f"数据源: {paths[0].parent}", font=small_font, fill="#555555")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output_path)


def pulse_width_from_header(header: DatHeader) -> float | str:
    for key, value in header.extra_info.items():
        if "ns" in str(key).lower() or "脉宽" in str(key):
            try:
                return float(value)
            except (TypeError, ValueError):
                return str(value)
    return ""


def write_summary_csv(path: Path, summaries: Iterable[ExportSummary]) -> None:
    rows = [asdict(item) for item in summaries]
    path.parent.mkdir(parents=True, exist_ok=True)
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=list(rows[0].keys()) if rows else list(ExportSummary.__annotations__.keys()))
        writer.writeheader()
        writer.writerows(rows)


def safe_filename(text: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", text).strip("_")


def process(args: argparse.Namespace) -> int:
    records = read_records(args.workbook)
    files = index_data_files(args.data_root)
    if not files:
        raise SystemExit(f"未找到可解析时间的数据文件: {args.data_root}")
    summaries: list[ExportSummary] = []
    args.output_dir.mkdir(parents=True, exist_ok=True)

    print(f"读取表格记录: {len(records)} 条")
    print(f"扫描数据文件: {len(files)} 个 ({args.data_root})")
    print(f"数据时间范围: {files[0].timestamp:%Y-%m-%d %H:%M:%S} ~ {files[-1].timestamp:%Y-%m-%d %H:%M:%S}")

    for index, record in enumerate(records, start=1):
        matched = files_for_record(files, record)
        summary = ExportSummary(
            group=record.group,
            seq=record.seq,
            start=record.start.isoformat(sep=" "),
            end=record.end.isoformat(sep=" "),
            operation=record.operation,
            flow_lh=record.flow_lh,
            file_count=len(matched),
            status="有数据" if matched else "无数据",
            first_file=matched[0].path.name if matched else "",
            last_file=matched[-1].path.name if matched else "",
        )
        print(f"[{index}/{len(records)}] {record.group} {record.seq}: {record.start:%H:%M:%S}-{record.end:%H:%M:%S}, 匹配 {len(matched)} 个文件")

        if matched and not args.match_only:
            output = args.output_dir / f"{safe_filename(record.label)}_phase_heatmap.png"
            try:
                display, header = load_display_matrix([item.path for item in matched], max_time_bins=args.max_time_bins)
                render_heatmap(display, header, record, [item.path for item in matched], output, vlim=args.vlim)
                summary.protocol = header.protocol
                summary.frequency_hz = header.frequency_hz
                summary.pulse_width_ns = pulse_width_from_header(header)
                summary.point_count = header.point_count
                summary.position_start_m = round(header.position_start_m, 3)
                summary.position_end_m = round(header.position_end_m, 3)
                summary.sample_count = header.sample_count
                summary.duration_s = round(header.time_span_s, 3)
                summary.image_path = str(output)
                print(f"  已输出: {output}")
            except Exception as exc:
                summary.status = "绘图失败"
                summary.message = str(exc)
                print(f"  绘图失败: {exc}")
        summaries.append(summary)

    csv_path = args.output_dir / "heatmap_summary.csv"
    json_path = args.output_dir / "heatmap_summary.json"
    write_summary_csv(csv_path, summaries)
    json_path.write_text(json.dumps([asdict(item) for item in summaries], ensure_ascii=False, indent=2), encoding="utf-8")
    print(f"SUMMARY_CSV {csv_path}")
    print(f"SUMMARY_JSON {json_path}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Export well-flow phase heatmaps from TDAS records.")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--data-root", type=Path, default=DEFAULT_DATA_ROOT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--vlim", type=float, default=DEFAULT_VLIM)
    parser.add_argument("--max-time-bins", type=int, default=DEFAULT_MAX_TIME_BINS)
    parser.add_argument("--match-only", action="store_true")
    args = parser.parse_args()
    return process(args)


if __name__ == "__main__":
    raise SystemExit(main())
