#!/usr/bin/env python
"""Batch the confirmed 2026-07-10 experiment groups for three DAS devices."""

from __future__ import annotations

import argparse
import bisect
import csv
import json
import math
from dataclasses import asdict, dataclass, replace
from datetime import date, datetime, time, timedelta
from pathlib import Path
from typing import Any

import numpy as np
from openpyxl import load_workbook
from PIL import Image, ImageDraw

from batch_well_flow_three_device_analysis import (
    Band,
    DatFile,
    DeviceConfig,
    PointRange,
    colorize_phase,
    compute_silixa_scale,
    find_font,
    index_data_files,
    parse_point_range,
    render_band_energy_heatmap,
    render_phase_heatmap,
    safe_filename,
    select_point_indices,
)
from tdas_dat_viewer import (
    DatHeader,
    SILIXA_NORMALIZED_MAX,
    make_stft_window,
    read_header,
    read_signal_matrix,
)


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = (
    PROJECT_ROOT
    / "井中流动分析"
    / "PSD分析_20260710"
    / "中国石油大学PSD分析法_实验记录与数据覆盖_已确认.xlsx"
)
DEFAULT_DATA_ROOT = Path(r"E:\北京试验")
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "井中流动分析" / "PSD分析_20260710" / "三设备分析结果"

DEVICE_CONFIGS = (
    DeviceConfig("silixia", "silixia", "silixia", 0),
    DeviceConfig("OV", "OV", "ov", 0),
    DeviceConfig("self_2u", "自研_2U", "自研_2U", 0),
)

BANDS = (
    Band("0~2Hz", 0.0, 2.0),
    Band("2~10Hz", 2.0, 10.0),
    Band("10~50Hz", 10.0, 50.0),
    Band("50~200Hz", 50.0, 200.0),
)

PHASE_COLORS = {
    "底噪": "#4b5563",
    "调泵": "#ef4444",
    "稳定": "#16a34a",
    "停止": "#111827",
}


@dataclass(frozen=True)
class PhaseInterval:
    label: str
    start: datetime
    end: datetime


@dataclass(frozen=True)
class GroupRecord:
    group: str
    date_text: str
    seq: str
    start: datetime
    end: datetime
    phases: tuple[PhaseInterval, ...]
    condition: str
    flow_lh: float
    unit: str
    device_ranges: dict[str, PointRange]

    @property
    def operation(self) -> str:
        return "/".join(phase.label for phase in self.phases)

    @property
    def flow_text(self) -> str:
        if self.flow_lh > 0:
            return f"{self.flow_lh:g}Lh"
        return safe_filename(self.condition) or "condition"

    @property
    def folder_name(self) -> str:
        return safe_filename(f"{self.seq}_{self.start:%H%M}-{self.end:%H%M}_{self.flow_text}")

    @property
    def title(self) -> str:
        condition = self.condition or (f"{self.flow_lh / 1000:g} m3/h" if self.flow_lh else "")
        return f"{self.group} {self.seq} | {self.start:%H:%M}-{self.end:%H:%M} | {condition}"


@dataclass(frozen=True)
class FileSegment:
    item: DatFile
    header: DatHeader
    start_index: int
    stop_index: int
    global_start_index: int

    @property
    def sample_count(self) -> int:
        return self.stop_index - self.start_index


@dataclass
class ResultRow:
    experiment_type: str
    experiment_id: str
    device: str
    start: str
    end: str
    phases: str
    condition: str
    requested_range_m: str
    actual_range_m: str = ""
    file_count: int = 0
    protocol: str = ""
    sample_rate_hz: int | str = ""
    point_count: int | str = ""
    position_start_m: float | str = ""
    position_end_m: float | str = ""
    selected_point_count: int = 0
    coverage_percent: float = 0.0
    missing_intervals: str = ""
    phase_image: str = ""
    band_images: dict[str, str] | None = None
    status: str = "pending"
    warnings: str = ""
    message: str = ""


def cell_text(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, time):
        return value.strftime("%H:%M:%S").removesuffix(":00")
    return str(value).strip()


def parse_date(value: Any) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    return date.fromisoformat(cell_text(value)[:10])


def parse_time(value: Any) -> time | None:
    text = cell_text(value)
    if not text:
        return None
    for fmt in ("%H:%M", "%H:%M:%S"):
        try:
            return datetime.strptime(text, fmt).time()
        except ValueError:
            continue
    raise ValueError(f"无法解析时间: {text}")


def read_records(workbook_path: Path) -> list[GroupRecord]:
    workbook = load_workbook(workbook_path, data_only=True, read_only=True)
    sheet = workbook["实验记录_已确认"]
    headers = {cell_text(cell.value): index for index, cell in enumerate(sheet[1])}

    def value(row: tuple[Any, ...], name: str) -> Any:
        return row[headers[name]]

    records: list[GroupRecord] = []
    for row_cells in sheet.iter_rows(min_row=2, values_only=True):
        if not row_cells or not row_cells[headers["实验编号"]]:
            continue
        test_date = parse_date(value(row_cells, "日期"))
        start_time = parse_time(value(row_cells, "实验开始时间"))
        end_time = parse_time(value(row_cells, "实验结束时间"))
        if start_time is None or end_time is None:
            raise ValueError(f"{value(row_cells, '实验编号')} 缺少实验组起止时间")
        start_dt = datetime.combine(test_date, start_time)
        end_dt = datetime.combine(test_date, end_time)
        phases: list[PhaseInterval] = []
        for label, start_name, end_name in (
            ("底噪", "底噪开始", "底噪结束"),
            ("调泵", "调泵开始", "调泵结束"),
            ("稳定", "稳定开始", "稳定结束"),
            ("停止", "停止开始", "停止结束"),
        ):
            phase_start = parse_time(value(row_cells, start_name))
            phase_end = parse_time(value(row_cells, end_name))
            if phase_start is None and phase_end is None:
                continue
            if phase_start is None or phase_end is None:
                raise ValueError(f"{value(row_cells, '实验编号')} 的{label}阶段时间不完整")
            phases.append(
                PhaseInterval(label, datetime.combine(test_date, phase_start), datetime.combine(test_date, phase_end))
            )
        condition = cell_text(value(row_cells, "气量/气水比说明"))
        flow_text = cell_text(value(row_cells, "液体流量(L/h)"))
        flow_lh = float(flow_text) if flow_text else 0.0
        if not condition and flow_lh:
            condition = f"{flow_lh / 1000:g} m3/h"
        records.append(
            GroupRecord(
                group=cell_text(value(row_cells, "实验类型")),
                date_text=test_date.isoformat(),
                seq=cell_text(value(row_cells, "实验编号")),
                start=start_dt,
                end=end_dt,
                phases=tuple(phases),
                condition=condition,
                flow_lh=flow_lh,
                unit="L/h",
                device_ranges={
                    "silixia": parse_point_range(value(row_cells, "silixia频谱点位范围(m)")),
                    "OV": parse_point_range(value(row_cells, "OV频谱点位范围(m)")),
                    "自研_2U": parse_point_range(value(row_cells, "自研_2U频谱点位范围(m)")),
                },
            )
        )
    return records


def validate_compatible(headers: list[DatHeader]) -> DatHeader:
    base = headers[0]
    for header in headers[1:]:
        checks = (
            (header.protocol == base.protocol, "协议"),
            (header.point_count == base.point_count, "点位数"),
            (header.frequency_hz == base.frequency_hz, "采样率"),
            (abs(header.position_start_m - base.position_start_m) <= 1e-6, "起始位置"),
            (abs(header.position_end_m - base.position_end_m) <= 1e-6, "结束位置"),
            (abs(header.point_interval_m - base.point_interval_m) <= 1e-6, "点位间距"),
        )
        for ok, label in checks:
            if not ok:
                raise ValueError(f"同一实验组内文件{label}不一致")
    return base


def build_segments(files: list[DatFile], record: GroupRecord) -> list[FileSegment]:
    timestamps = [item.timestamp for item in files]
    first = max(0, bisect.bisect_right(timestamps, record.start) - 1)
    candidates = files[first:]
    segments: list[FileSegment] = []
    stitched_samples = 0
    target_samples: int | None = None
    for item in candidates:
        header = read_header(item.path)
        rate = max(1, int(header.frequency_hz))
        if target_samples is None:
            target_samples = int(round((record.end - record.start).total_seconds() * rate))
        file_end = item.timestamp + timedelta(seconds=header.sample_count / rate)
        if file_end <= record.start:
            continue
        remaining = max(0, target_samples - stitched_samples)
        if remaining <= 0:
            break
        if not segments:
            local_start = max(0, int(math.ceil((record.start - item.timestamp).total_seconds() * rate - 1e-9)))
        else:
            local_start = 0
        local_stop = min(header.sample_count, local_start + remaining)
        if local_stop <= local_start:
            continue
        segments.append(FileSegment(item, header, local_start, local_stop, stitched_samples))
        stitched_samples += local_stop - local_start
        if target_samples is not None and stitched_samples >= target_samples:
            break
    if not segments:
        return []
    validate_compatible([segment.header for segment in segments])
    return segments


def stitched_sample_count(segments: list[FileSegment]) -> int:
    if not segments:
        return 0
    return max(segment.global_start_index + segment.sample_count for segment in segments)


def load_group_phase_display(
    segments: list[FileSegment], record: GroupRecord, max_time_bins: int
) -> tuple[np.ndarray, DatHeader, float]:
    base = segments[0].header
    rate = max(1, int(base.frequency_hz))
    total_samples = stitched_sample_count(segments)
    bin_count = max(1, min(max_time_bins, total_samples))
    edges = np.linspace(0, total_samples, bin_count + 1, dtype=np.int64)
    display = np.zeros((base.point_count, bin_count), dtype=np.float32)
    best_abs = np.full((base.point_count, bin_count), -np.inf, dtype=np.float32)
    rows = np.arange(base.point_count)
    raw_max_abs = 0.0

    for segment in segments:
        matrix = read_signal_matrix(segment.item.path, segment.header).astype(np.float32, copy=False)
        matrix = matrix[:, segment.start_index : segment.stop_index]
        if matrix.size:
            raw_max_abs = max(raw_max_abs, float(np.max(np.abs(matrix))))
        segment_start = segment.global_start_index
        segment_stop = min(total_samples, segment_start + matrix.shape[1])
        matrix = matrix[:, : max(0, segment_stop - segment_start)]
        if matrix.shape[1] == 0:
            continue
        first_bin = max(0, int(np.searchsorted(edges, segment_start, side="right") - 1))
        last_bin = min(bin_count - 1, int(np.searchsorted(edges, segment_stop - 1, side="right") - 1))
        for bin_index in range(first_bin, last_bin + 1):
            overlap_start = max(int(edges[bin_index]), segment_start)
            overlap_stop = min(int(edges[bin_index + 1]), segment_stop)
            if overlap_stop <= overlap_start:
                continue
            local_start = overlap_start - segment_start
            local_stop = overlap_stop - segment_start
            block = matrix[:, local_start:local_stop]
            indices = np.argmax(np.abs(block), axis=1)
            values = block[rows, indices]
            abs_values = np.abs(values)
            update = abs_values > best_abs[:, bin_index]
            display[update, bin_index] = values[update]
            best_abs[update, bin_index] = abs_values[update]

    scale = 1.0
    if base.protocol == "Silixa TDMS" and raw_max_abs > 0.0:
        scale = float(SILIXA_NORMALIZED_MAX / raw_max_abs)
        display *= scale
    combined = replace(
        base,
        sample_count=total_samples,
        acquisition_time_s=total_samples / rate,
        timestamp_text=record.start.isoformat(sep=" "),
    )
    return display, combined, scale


def coverage_gaps(segments: list[FileSegment], record: GroupRecord) -> tuple[list[tuple[float, float]], float]:
    rate = max(1, int(segments[0].header.frequency_hz))
    total_samples = int(round((record.end - record.start).total_seconds() * rate))
    intervals = sorted(
        (
            max(0, segment.global_start_index),
            min(total_samples, segment.global_start_index + segment.sample_count),
        )
        for segment in segments
    )
    merged: list[list[int]] = []
    for start, end in intervals:
        if end <= start:
            continue
        if not merged or start > merged[-1][1] + 1:
            merged.append([start, end])
        else:
            merged[-1][1] = max(merged[-1][1], end)
    stitched = stitched_sample_count(segments)
    gaps = [] if stitched >= total_samples else [(stitched / rate, total_samples / rate)]
    coverage_percent = 100.0 * min(stitched, total_samples) / max(1, total_samples)
    return gaps, coverage_percent


def format_missing_intervals(record: GroupRecord, gaps: list[tuple[float, float]]) -> str:
    parts = []
    for start_s, end_s in gaps:
        start = record.start + timedelta(seconds=start_s)
        end = record.start + timedelta(seconds=end_s)
        parts.append(f"{start:%H:%M:%S}-{end:%H:%M:%S}")
    return "; ".join(parts)


def compute_group_band_energy(
    segments: list[FileSegment],
    point_indices: np.ndarray,
    bands: tuple[Band, ...],
    scale: float,
    window_s: float,
    overlap_percent: float,
    total_duration_s: float,
) -> tuple[np.ndarray, np.ndarray]:
    base = segments[0].header
    sample_rate_hz = max(1, int(base.frequency_hz))
    window_samples = int(round(window_s * sample_rate_hz))
    hop_samples = max(1, int(round(window_samples * (1.0 - overlap_percent / 100.0))))
    freqs = np.fft.rfftfreq(window_samples, d=1.0 / sample_rate_hz).astype(np.float32)
    df = float(freqs[1] - freqs[0])
    masks = []
    for band in bands:
        if band.high_hz > float(freqs[-1]):
            raise ValueError(f"频带{band.label}超过Nyquist频率{float(freqs[-1]):.1f}Hz")
        mask = np.flatnonzero((freqs >= band.low_hz) & (freqs <= band.high_hz))
        if mask.size == 0:
            raise ValueError(f"频带{band.label}没有FFT频点")
        masks.append(mask)

    window = make_stft_window("Hann", window_samples).astype(np.float32, copy=False)
    window_power = float(np.sum(window * window))
    buffer = np.empty((int(point_indices.size), 0), dtype=np.float32)
    buffer_start = 0
    loaded_until = 0
    next_start = 0
    initialized = False
    columns: list[np.ndarray] = []
    centers: list[float] = []

    for segment in segments:
        matrix = read_signal_matrix(segment.item.path, segment.header)[point_indices, :].astype(np.float32, copy=False)
        matrix = matrix[:, segment.start_index : segment.stop_index]
        if scale != 1.0:
            matrix = matrix * scale
        global_start = segment.global_start_index
        if not initialized or global_start > loaded_until + 1:
            buffer = matrix
            buffer_start = global_start
            loaded_until = global_start + matrix.shape[1]
            next_start = global_start
            initialized = True
        else:
            if global_start < loaded_until:
                overlap = loaded_until - global_start
                if overlap >= matrix.shape[1]:
                    continue
                matrix = matrix[:, overlap:]
                global_start += overlap
            buffer = np.concatenate((buffer, matrix), axis=1)
            loaded_until = global_start + matrix.shape[1]

        while next_start + window_samples <= loaded_until:
            local_start = next_start - buffer_start
            if local_start < 0:
                next_start = buffer_start
                local_start = 0
            local_stop = local_start + window_samples
            if local_stop > buffer.shape[1]:
                break
            block = buffer[:, local_start:local_stop]
            block = block - np.mean(block, axis=1, keepdims=True)
            spectrum = np.fft.rfft(block * window[np.newaxis, :], axis=1)
            psd = (np.abs(spectrum) ** 2) / (sample_rate_hz * window_power)
            if psd.shape[1] > 1:
                if window_samples % 2 == 0:
                    psd[:, 1:-1] *= 2.0
                else:
                    psd[:, 1:] *= 2.0
            band_values = np.empty((len(bands), int(point_indices.size)), dtype=np.float32)
            for band_index, mask in enumerate(masks):
                energy = np.sum(psd[:, mask], axis=1) * df
                band_values[band_index] = 10.0 * np.log10(np.maximum(energy, 1e-24))
            columns.append(band_values)
            centers.append((next_start + window_samples / 2.0) / sample_rate_hz)
            next_start += hop_samples

        if next_start > buffer_start:
            drop = min(next_start - buffer_start, buffer.shape[1])
            buffer = buffer[:, drop:]
            buffer_start += drop

    if not columns:
        raise ValueError("完整实验组内没有生成STFT时间窗")
    raw_energy = np.stack(columns, axis=2)
    total_samples = int(round(total_duration_s * sample_rate_hz))
    grid_starts = np.arange(0, max(0, total_samples - window_samples) + 1, hop_samples, dtype=np.int64)
    grid_centers = (grid_starts + window_samples / 2.0) / sample_rate_hz
    energy = np.full(
        (len(bands), int(point_indices.size), int(grid_centers.size)),
        np.nan,
        dtype=np.float32,
    )
    for column_index, center_s in enumerate(centers):
        start_sample = int(round(center_s * sample_rate_hz - window_samples / 2.0))
        grid_index = int(round(start_sample / hop_samples))
        if 0 <= grid_index < grid_centers.size:
            energy[:, :, grid_index] = raw_energy[:, :, column_index]
    return energy, np.asarray(grid_centers, dtype=np.float32)


def annotate_phase_heatmap(path: Path, record: GroupRecord, gaps: list[tuple[float, float]]) -> None:
    image = Image.open(path).convert("RGBA")
    draw = ImageDraw.Draw(image)
    left, top, plot_w, plot_h = 145, 165, 1360, 760
    right = left + plot_w
    duration = max((record.end - record.start).total_seconds(), 1.0)
    font = find_font(19)
    overlay = Image.new("RGBA", image.size, (0, 0, 0, 0))
    overlay_draw = ImageDraw.Draw(overlay)
    gap_font = find_font(22)
    for gap_start, gap_end in gaps:
        y1 = top + int(round(max(0.0, gap_start) / duration * plot_h))
        y2 = top + int(round(min(duration, gap_end) / duration * plot_h))
        if y2 <= y1:
            continue
        overlay_draw.rectangle((left, y1, right, y2), fill=(75, 85, 99, 175))
        label = "无数据"
        bbox = overlay_draw.textbbox((0, 0), label, font=gap_font)
        label_w = bbox[2] - bbox[0]
        overlay_draw.text(((left + right - label_w) // 2, (y1 + y2) // 2 - 12), label, font=gap_font, fill=(255, 255, 255, 255))
    image = Image.alpha_composite(image, overlay)
    draw = ImageDraw.Draw(image)
    for phase in record.phases:
        start_s = (phase.start - record.start).total_seconds()
        end_s = (phase.end - record.start).total_seconds()
        for value in (start_s, end_s):
            if 0.0 < value < duration:
                y = top + int(round(value / duration * plot_h))
                draw.line((left, y, right, y), fill="#ffffff", width=3)
                draw.line((left, y + 2, right, y + 2), fill=PHASE_COLORS[phase.label], width=2)
        mid_y = top + int(round(((start_s + end_s) / 2.0) / duration * plot_h))
        draw.text((right + 8, mid_y - 10), phase.label, font=font, fill=PHASE_COLORS[phase.label])
    image.convert("RGB").save(path)


def annotate_band_heatmap(
    path: Path,
    record: GroupRecord,
    centers_s: np.ndarray,
    gaps: list[tuple[float, float]],
) -> None:
    image = Image.open(path).convert("RGBA")
    overlay = Image.new("RGBA", image.size, (0, 0, 0, 0))
    overlay_draw = ImageDraw.Draw(overlay)
    left, top, plot_w, plot_h = 160, 170, 1345, 750
    x_min = float(centers_s[0])
    x_max = float(centers_s[-1]) if centers_s.size > 1 else x_min + 1.0
    font = find_font(18)
    gap_font = find_font(22)
    for gap_start, gap_end in gaps:
        clipped_start = max(gap_start, x_min)
        clipped_end = min(gap_end, x_max)
        if clipped_end <= clipped_start:
            continue
        x1 = left + int(round((clipped_start - x_min) / max(x_max - x_min, 1e-9) * plot_w))
        x2 = left + int(round((clipped_end - x_min) / max(x_max - x_min, 1e-9) * plot_w))
        overlay_draw.rectangle((x1, top, x2, top + plot_h), fill=(75, 85, 99, 175))
        label = "无数据"
        bbox = overlay_draw.textbbox((0, 0), label, font=gap_font)
        label_w = bbox[2] - bbox[0]
        if x2 - x1 > label_w + 12:
            overlay_draw.text(((x1 + x2 - label_w) // 2, top + plot_h // 2 - 12), label, font=gap_font, fill=(255, 255, 255, 255))
    image = Image.alpha_composite(image, overlay)
    draw = ImageDraw.Draw(image)
    for phase in record.phases:
        start_s = (phase.start - record.start).total_seconds()
        end_s = (phase.end - record.start).total_seconds()
        for value in (start_s, end_s):
            if x_min < value < x_max:
                x = left + int(round((value - x_min) / max(x_max - x_min, 1e-9) * plot_w))
                draw.line((x, top, x, top + plot_h), fill="#ffffff", width=3)
                draw.line((x + 2, top, x + 2, top + plot_h), fill=PHASE_COLORS[phase.label], width=2)
        clipped_start = max(start_s, x_min)
        clipped_end = min(end_s, x_max)
        if clipped_end > clipped_start:
            mid = (clipped_start + clipped_end) / 2.0
            x = left + int(round((mid - x_min) / max(x_max - x_min, 1e-9) * plot_w))
            bbox = draw.textbbox((0, 0), phase.label, font=font)
            width = bbox[2] - bbox[0]
            draw.rectangle((x - width // 2 - 5, top + 8, x + width // 2 + 5, top + 33), fill="#ffffff")
            draw.text((x - width // 2, top + 10), phase.label, font=font, fill=PHASE_COLORS[phase.label])
    image.convert("RGB").save(path)


def process_record_device(
    record: GroupRecord,
    device: DeviceConfig,
    files: list[DatFile],
    output_dir: Path,
    max_time_bins: int,
    phase_vlim: float,
    window_s: float,
    overlap_percent: float,
    match_only: bool,
) -> ResultRow:
    requested_range = record.device_ranges[device.title]
    result = ResultRow(
        experiment_type=record.group,
        experiment_id=record.seq,
        device=device.title,
        start=record.start.isoformat(sep=" "),
        end=record.end.isoformat(sep=" "),
        phases=record.operation,
        condition=record.condition,
        requested_range_m=requested_range.text,
        band_images={},
    )
    try:
        segments = build_segments(files, record)
        result.file_count = len(segments)
        if not segments:
            result.status = "no_data"
            result.message = "完整实验组时间内没有可用数据"
            return result
        gaps, coverage_percent = coverage_gaps(segments, record)
        missing_text = format_missing_intervals(record, gaps)
        result.coverage_percent = round(coverage_percent, 2)
        result.missing_intervals = missing_text
        base = segments[0].header
        indices, positions, range_warning = select_point_indices(base, requested_range)
        result.protocol = base.protocol
        result.sample_rate_hz = base.frequency_hz
        result.point_count = base.point_count
        result.position_start_m = round(base.position_start_m, 3)
        result.position_end_m = round(base.position_end_m, 3)
        result.selected_point_count = int(indices.size)
        result.actual_range_m = f"{float(positions[0]):.3f}-{float(positions[-1]):.3f}"
        result.warnings = range_warning
        if match_only:
            result.status = "matched"
            return result

        device_dir = output_dir / safe_filename(record.group) / record.folder_name / safe_filename(device.title)
        phase_path = device_dir / "phase" / f"{record.folder_name}_{safe_filename(device.title)}_完整实验组_相位热力图.png"
        display, combined_header, scale = load_group_phase_display(segments, record, max_time_bins)
        render_phase_heatmap(
            display,
            combined_header,
            record,
            device,
            [segment.item.path for segment in segments],
            phase_path,
            vlim=phase_vlim,
        )
        result.phase_image = str(phase_path)

        energy_by_band, centers_s = compute_group_band_energy(
            segments,
            indices,
            BANDS,
            scale,
            window_s,
            overlap_percent,
            stitched_sample_count(segments) / max(1, int(base.frequency_hz)),
        )
        band_dir = device_dir / "band_energy"
        for band_index, band in enumerate(BANDS):
            output_path = band_dir / f"{record.folder_name}_{safe_filename(device.title)}_{band.filename_label}_PSD频带能量图.png"
            coverage_warning = f"数据覆盖 {coverage_percent:.1f}%"
            if missing_text:
                coverage_warning += f"；缺失 {missing_text}"
            combined_warning = "；".join(text for text in (range_warning, coverage_warning) if text)
            render_band_energy_heatmap(
                energy_by_band[band_index],
                positions,
                centers_s,
                band,
                record,
                device,
                requested_range,
                output_path,
                warnings=combined_warning,
                robust_levels=True,
            )
            result.band_images[band.label] = str(output_path)
        result.status = "partial" if gaps else "ok"
    except Exception as exc:
        result.status = "failed"
        result.message = str(exc)
    return result


def write_outputs(output_dir: Path, results: list[ResultRow]) -> None:
    output_dir.mkdir(parents=True, exist_ok=True)
    json_path = output_dir / "PSD_20260710_三设备分析清单.json"
    json_path.write_text(json.dumps([asdict(row) for row in results], ensure_ascii=False, indent=2), encoding="utf-8")
    csv_path = output_dir / "PSD_20260710_三设备分析清单.csv"
    csv_rows = []
    for row in results:
        item = asdict(row)
        item["band_images"] = json.dumps(item["band_images"], ensure_ascii=False)
        csv_rows.append(item)
    with csv_path.open("w", encoding="utf-8-sig", newline="") as handle:
        writer = csv.DictWriter(handle, fieldnames=list(csv_rows[0].keys()))
        writer.writeheader()
        writer.writerows(csv_rows)


def main() -> int:
    parser = argparse.ArgumentParser(description="中国石油大学PSD分析法：2026-07-10三设备完整实验组批处理")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--data-root", type=Path, default=DEFAULT_DATA_ROOT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--device", default="")
    parser.add_argument("--record", default="")
    parser.add_argument("--match-only", action="store_true")
    parser.add_argument("--phase-vlim", type=float, default=0.15)
    parser.add_argument("--max-time-bins", type=int, default=4096)
    parser.add_argument("--stft-window-s", type=float, default=2.0)
    parser.add_argument("--stft-overlap", type=float, default=50.0)
    args = parser.parse_args()

    records = read_records(args.workbook)
    if args.record:
        records = [record for record in records if record.seq.lower() == args.record.lower()]
    devices = [
        device
        for device in DEVICE_CONFIGS
        if not args.device
        or args.device.lower() in {device.key.lower(), device.title.lower(), device.data_subdir.lower()}
    ]
    if not records:
        raise SystemExit(f"未找到实验组: {args.record}")
    if not devices:
        raise SystemExit(f"未找到设备: {args.device}")

    files_by_device = {}
    for device in devices:
        files = index_data_files(args.data_root / device.data_subdir)
        files_by_device[device.title] = files
        if files:
            print(f"{device.title}: {len(files)} files, {files[0].timestamp} ~ {files[-1].timestamp}", flush=True)
        else:
            print(f"{device.title}: no files", flush=True)

    results: list[ResultRow] = []
    total = len(records) * len(devices)
    for index, record in enumerate(records):
        for device_index, device in enumerate(devices):
            count = index * len(devices) + device_index + 1
            print(f"[{count}/{total}] {record.group} {record.seq} {device.title} {record.start:%H:%M}-{record.end:%H:%M}", flush=True)
            result = process_record_device(
                record,
                device,
                files_by_device[device.title],
                args.output_dir,
                args.max_time_bins,
                args.phase_vlim,
                args.stft_window_s,
                args.stft_overlap,
                args.match_only,
            )
            results.append(result)
            print(f"  {result.status}: files={result.file_count}, points={result.selected_point_count}, {result.message}", flush=True)
    write_outputs(args.output_dir, results)
    failed = [row for row in results if row.status not in {"ok", "partial", "matched"}]
    partial = [row for row in results if row.status == "partial"]
    print(f"RESULTS total={len(results)} partial={len(partial)} failed={len(failed)} output={args.output_dir}", flush=True)
    return 1 if failed else 0


if __name__ == "__main__":
    raise SystemExit(main())
