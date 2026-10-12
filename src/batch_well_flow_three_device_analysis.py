#!/usr/bin/env python
"""Batch TDAS well-flow analysis for three devices.

The data readers and STFT window implementation are reused from
``tdas_dat_viewer.py``. This script adds the experiment-table parsing,
three-device matching, point-range clipping, phase heatmap rendering, and
band-energy heatmap rendering needed for the well-flow records.
"""

from __future__ import annotations

import argparse
import csv
import json
import math
import re
import sys
import types
from dataclasses import asdict, dataclass, field
from datetime import date, datetime, time
from pathlib import Path
from typing import Iterable

import numpy as np
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = PROJECT_ROOT / "井中流动分析" / "井中流动分析_两组实验记录.xlsx"
DEFAULT_DATA_ROOT = Path(r"E:\北京试验")
DEFAULT_OUTPUT_DIR = PROJECT_ROOT / "井中流动分析" / "三设备分析结果"
DEFAULT_PHASE_VLIM = 0.3
DEFAULT_MAX_TIME_BINS = 4096
DEFAULT_STFT_WINDOW_S = 2.0
DEFAULT_STFT_OVERLAP_PERCENT = 50.0


def install_gui_stubs() -> None:
    """Let tdas_dat_viewer import in a non-GUI batch process."""
    class DummySignal:
        def connect(self, *_args: object, **_kwargs: object) -> None:
            return None

        def emit(self, *_args: object, **_kwargs: object) -> None:
            return None

    class DummyBase:
        def __init__(self, *_args: object, **_kwargs: object) -> None:
            pass

        def __call__(self, *_args: object, **_kwargs: object) -> "DummyBase":
            return self

        def __getattr__(self, _name: str) -> "DummyBase":
            return self

    class DummyModule(types.ModuleType):
        def __getattr__(self, name: str) -> object:
            if name == "pyqtSignal":
                return lambda *_args, **_kwargs: DummySignal()
            if name == "pyqtSlot":
                return lambda *_args, **_kwargs: (lambda func: func)
            if name in {"QObject", "QRunnable", "QMainWindow", "QWidget", "QDialog"}:
                return DummyBase
            return DummyBase

    if "PyQt6" not in sys.modules:
        pyqt6 = types.ModuleType("PyQt6")
        qtcore = DummyModule("PyQt6.QtCore")
        qtgui = DummyModule("PyQt6.QtGui")
        qtwidgets = DummyModule("PyQt6.QtWidgets")
        pyqt6.QtCore = qtcore  # type: ignore[attr-defined]
        pyqt6.QtGui = qtgui  # type: ignore[attr-defined]
        pyqt6.QtWidgets = qtwidgets  # type: ignore[attr-defined]
        sys.modules["PyQt6"] = pyqt6
        sys.modules["PyQt6.QtCore"] = qtcore
        sys.modules["PyQt6.QtGui"] = qtgui
        sys.modules["PyQt6.QtWidgets"] = qtwidgets
    if "pyqtgraph" not in sys.modules:
        pg = DummyModule("pyqtgraph")
        pg.opengl = DummyModule("pyqtgraph.opengl")  # type: ignore[attr-defined]
        sys.modules["pyqtgraph"] = pg
        sys.modules["pyqtgraph.opengl"] = pg.opengl  # type: ignore[attr-defined]


install_gui_stubs()
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from tdas_dat_viewer import (  # noqa: E402
    DatHeader,
    SILIXA_NORMALIZED_MAX,
    make_stft_window,
    read_header,
    read_signal_matrix,
)


@dataclass(frozen=True)
class DeviceConfig:
    key: str
    title: str
    data_subdir: str
    range_column_index: int


DEVICE_CONFIGS: tuple[DeviceConfig, ...] = (
    DeviceConfig("silixia", "silixia", "silixia", 8),
    DeviceConfig("OV", "OV", "ov", 9),
    DeviceConfig("self_2u", "自研_2U", "自研_2U", 10),
)


@dataclass(frozen=True)
class Band:
    label: str
    low_hz: float
    high_hz: float
    zero_only: bool = False

    @property
    def filename_label(self) -> str:
        return re.sub(r"[^0-9A-Za-z]+", "_", self.label).strip("_")


@dataclass(frozen=True)
class DatFile:
    path: Path
    timestamp: datetime


@dataclass(frozen=True)
class PointRange:
    text: str
    low_m: float
    high_m: float


@dataclass(frozen=True)
class ExperimentRecord:
    group: str
    date_text: str
    seq: str
    start: datetime
    middle_text: str
    end: datetime
    end_text: str
    operation: str
    flow_lh: float
    unit: str
    device_ranges: dict[str, PointRange]

    @property
    def flow_text(self) -> str:
        return f"{self.flow_lh:g}{self.unit.replace('/', '')}"

    @property
    def folder_name(self) -> str:
        return safe_filename(f"{self.seq}_{self.start:%H%M%S}-{self.end:%H%M%S}_{self.flow_text}")

    @property
    def title(self) -> str:
        return f"{self.group} {self.seq} | {self.start:%H:%M:%S}-{self.end:%H:%M:%S} | {self.operation} {self.flow_lh:g}{self.unit}"


@dataclass
class DeviceRecordSummary:
    group: str
    seq: str
    device: str
    start: str
    end: str
    operation: str
    flow_lh: float
    requested_range_m: str
    actual_range_m: str = ""
    selected_point_count: int = 0
    file_count: int = 0
    status: str = "no_data"
    protocol: str = ""
    frequency_hz: int | str = ""
    point_count: int | str = ""
    position_start_m: float | str = ""
    position_end_m: float | str = ""
    sample_count: int | str = ""
    duration_s: float | str = ""
    first_file: str = ""
    last_file: str = ""
    phase_image: str = ""
    band_image_count: int = 0
    output_folder: str = ""
    message: str = ""
    warnings: str = ""


def as_text(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, time):
        return f"{value.hour:02d}:{value.minute:02d}:{value.second:02d}".removesuffix(":00")
    return str(value).strip()


def parse_date(value: object) -> date:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    text = as_text(value)
    match = re.search(r"\d{4}-\d{2}-\d{2}", text)
    if not match:
        raise ValueError(f"无法解析日期: {text}")
    return date.fromisoformat(match.group(0))


def parse_time_value(value: object, use_last: bool = False) -> time:
    if isinstance(value, datetime):
        return value.time().replace(microsecond=0)
    if isinstance(value, time):
        return value.replace(microsecond=0)
    text = as_text(value)
    matches = re.findall(r"(\d{1,2}):(\d{2})(?::(\d{2}))?", text)
    if not matches:
        raise ValueError(f"无法解析时间: {text}")
    hour, minute, second = matches[-1 if use_last else 0]
    return time(int(hour), int(minute), int(second or 0))


def parse_point_range(value: object) -> PointRange:
    text = as_text(value)
    numbers = re.findall(r"-?\d+(?:\.\d+)?", text)
    if len(numbers) >= 2:
        low, high = float(numbers[0]), float(numbers[1])
    elif len(numbers) == 1:
        low = high = float(numbers[0])
    else:
        raise ValueError(f"无法解析点位范围: {text}")
    if high < low:
        low, high = high, low
    return PointRange(text=text, low_m=low, high_m=high)


def read_records(workbook_path: Path) -> list[ExperimentRecord]:
    workbook = load_workbook(workbook_path, data_only=True)
    if len(workbook.worksheets) < 3:
        raise ValueError(f"工作簿至少需要包含参数表、第一组实验、第二组实验: {workbook_path}")

    sheet_groups = (
        (workbook.worksheets[1], "第一组实验"),
        (workbook.worksheets[2], "第二组实验"),
    )
    records: list[ExperimentRecord] = []
    for sheet, group in sheet_groups:
        for row in sheet.iter_rows(min_row=4, values_only=True):
            if not row or not row[0]:
                continue
            test_date = parse_date(row[0])
            start_time = parse_time_value(row[2])
            middle_text = as_text(row[3])
            end_time = parse_time_value(row[4], use_last=True)
            start_dt = datetime.combine(test_date, start_time)
            end_dt = datetime.combine(test_date, end_time)
            if end_dt <= start_dt:
                raise ValueError(f"{group} {row[1]} 结束时间不晚于开始时间")
            device_ranges = {
                config.title: parse_point_range(row[config.range_column_index])
                for config in DEVICE_CONFIGS
            }
            records.append(
                ExperimentRecord(
                    group=group,
                    date_text=as_text(row[0]),
                    seq=as_text(row[1]),
                    start=start_dt,
                    middle_text=middle_text,
                    end=end_dt,
                    end_text=as_text(row[4]),
                    operation=as_text(row[5]),
                    flow_lh=float(row[6]),
                    unit=as_text(row[7]) or "L/h",
                    device_ranges=device_ranges,
                )
            )
    return apply_user_confirmed_times(records)


def apply_user_confirmed_times(records: list[ExperimentRecord]) -> list[ExperimentRecord]:
    corrected: list[ExperimentRecord] = []
    for record in records:
        if record.seq == "3" and int(record.flow_lh) == 150:
            corrected.append(
                ExperimentRecord(
                    **{
                        **record.__dict__,
                        "middle_text": "15:56",
                        "end": datetime.combine(record.start.date(), time(16, 1)),
                        "end_text": "15:56-16:01",
                    }
                )
            )
            continue
        if record.group == "第一组实验" and record.seq == "1" and int(record.flow_lh) == 260:
            corrected.append(
                ExperimentRecord(
                    **{
                        **record.__dict__,
                        "start": datetime.combine(record.start.date(), time(15, 27)),
                        "middle_text": "15:38",
                        "end": datetime.combine(record.start.date(), time(15, 43)),
                        "end_text": "15:43",
                    }
                )
            )
            continue
        if record.group == "第二组实验" and record.seq.upper() == "E" and int(record.flow_lh) == 250:
            corrected.append(
                ExperimentRecord(
                    **{
                        **record.__dict__,
                        "start": datetime.combine(record.start.date(), time(17, 5)),
                        "middle_text": "17:15",
                        "end": datetime.combine(record.start.date(), time(17, 20)),
                        "end_text": "17:15-17:20",
                    }
                )
            )
            continue
        corrected.append(record)
    return corrected


def parse_file_timestamp(path: Path) -> datetime | None:
    name = path.name
    patterns = [
        (r"(\d{4})-(\d{2})-(\d{2})-(\d{2})-(\d{2})-(\d{2})-out\.dat$", "%Y-%m-%d-%H-%M-%S"),
        (r"(\d{14})\.\d+\.dat$", "%Y%m%d%H%M%S"),
        (r"_(\d{8})_(\d{6})(?:\.|_)", "%Y%m%d%H%M%S"),
    ]
    for pattern, fmt in patterns:
        match = re.search(pattern, name, re.IGNORECASE)
        if not match:
            continue
        if len(match.groups()) == 1:
            stamp = match.group(1)
        elif fmt == "%Y%m%d%H%M%S":
            stamp = match.group(1) + match.group(2)
        else:
            stamp = "-".join(match.groups())
        try:
            return datetime.strptime(stamp, fmt)
        except ValueError:
            continue
    return None


def index_data_files(data_root: Path) -> list[DatFile]:
    patterns = ("*.dat", "*.DAT", "*.tdms", "*.TDMS")
    files: list[DatFile] = []
    seen_paths: set[str] = set()
    for pattern in patterns:
        for path in data_root.rglob(pattern):
            path_key = str(path.resolve()).casefold()
            if path_key in seen_paths:
                continue
            seen_paths.add(path_key)
            timestamp = parse_file_timestamp(path)
            if timestamp is not None:
                files.append(DatFile(path=path, timestamp=timestamp))
    return sorted(files, key=lambda item: item.timestamp)


def files_for_record(files: list[DatFile], record: ExperimentRecord) -> list[DatFile]:
    return [item for item in files if record.start <= item.timestamp <= record.end]


def validate_headers(paths: list[Path]) -> tuple[DatHeader, dict[Path, DatHeader]]:
    headers = {path: read_header(path) for path in paths}
    base = headers[paths[0]]
    total_samples = 0
    for path in paths:
        header = headers[path]
        if header.protocol != base.protocol:
            raise ValueError(f"协议类型不同: {path.name}={header.protocol}, {paths[0].name}={base.protocol}")
        if header.point_count != base.point_count:
            raise ValueError(f"点位数不同: {path.name}={header.point_count}, {paths[0].name}={base.point_count}")
        if abs(header.position_start_m - base.position_start_m) > 1e-6:
            raise ValueError(f"起始位置不同: {path.name}={header.position_start_m}, {paths[0].name}={base.position_start_m}")
        if abs(header.position_end_m - base.position_end_m) > 1e-6:
            raise ValueError(f"结束位置不同: {path.name}={header.position_end_m}, {paths[0].name}={base.position_end_m}")
        if abs(header.point_interval_m - base.point_interval_m) > 1e-6:
            raise ValueError(f"点位间隔不同: {path.name}={header.point_interval_m}, {paths[0].name}={base.point_interval_m}")
        if header.frequency_hz != base.frequency_hz:
            raise ValueError(f"采样率不同: {path.name}={header.frequency_hz}, {paths[0].name}={base.frequency_hz}")
        if header.protocol == "CDAS" and header.channel != base.channel:
            raise ValueError(f"通道号不同: {path.name}=通道{header.channel}, {paths[0].name}=通道{base.channel}")
        total_samples += header.sample_count

    combined = DatHeader(
        protocol=base.protocol,
        system_id=base.system_id,
        year=base.year,
        month=base.month,
        day=base.day,
        hour=base.hour,
        minute=base.minute,
        second=base.second,
        length=base.length,
        global_alarm=base.global_alarm,
        acquisition_status=base.acquisition_status,
        channel=base.channel,
        frequency_hz=base.frequency_hz,
        acquisition_time_s=total_samples / max(1, base.frequency_hz),
        point_interval_m=base.point_interval_m,
        point_count=base.point_count,
        sample_count=total_samples,
        position_start_m=base.position_start_m,
        raw_value_note=base.raw_value_note,
        header_size=base.header_size,
        bytes_per_sample=base.bytes_per_sample,
        timestamp_text=base.timestamp_text,
        extra_info=base.extra_info,
    )
    return combined, headers


def compute_silixa_scale(paths: list[Path], headers: dict[Path, DatHeader]) -> float:
    max_abs = 0.0
    for path in paths:
        matrix = read_signal_matrix(path, headers[path])
        if matrix.size:
            max_abs = max(max_abs, float(np.max(np.abs(matrix))))
    if max_abs <= 0.0:
        return 1.0
    return float(SILIXA_NORMALIZED_MAX / max_abs)


def load_phase_display_matrix(paths: list[Path], max_time_bins: int) -> tuple[np.ndarray, DatHeader]:
    header, headers = validate_headers(paths)
    bin_count = max(1, min(int(max_time_bins), int(header.sample_count)))
    edges = np.linspace(0, header.sample_count, bin_count + 1, dtype=np.int64)
    display = np.zeros((header.point_count, bin_count), dtype=np.float32)
    best_abs = np.full((header.point_count, bin_count), -np.inf, dtype=np.float32)
    row_indices = np.arange(header.point_count)
    scale = compute_silixa_scale(paths, headers) if header.protocol == "Silixa TDMS" else 1.0

    sample_offset = 0
    for path in paths:
        file_header = headers[path]
        matrix = read_signal_matrix(path, file_header).astype(np.float32, copy=False)
        if scale != 1.0:
            matrix = matrix * scale
        file_start = sample_offset
        file_stop = sample_offset + file_header.sample_count
        first_bin = max(0, int(np.searchsorted(edges, file_start, side="right") - 1))
        last_bin = min(bin_count - 1, int(np.searchsorted(edges, file_stop - 1, side="right") - 1))
        for bin_index in range(first_bin, last_bin + 1):
            global_start = int(edges[bin_index])
            global_stop = int(edges[bin_index + 1])
            overlap_start = max(global_start, file_start)
            overlap_stop = min(global_stop, file_stop)
            if overlap_stop <= overlap_start:
                continue
            local_start = overlap_start - file_start
            local_stop = overlap_stop - file_start
            block = matrix[:, local_start:local_stop]
            max_indices = np.argmax(np.abs(block), axis=1)
            values = block[row_indices, max_indices]
            abs_values = np.abs(values)
            update_mask = abs_values > best_abs[:, bin_index]
            display[update_mask, bin_index] = values[update_mask]
            best_abs[update_mask, bin_index] = abs_values[update_mask]
        sample_offset = file_stop
    return display, header


def positions_for_header(header: DatHeader) -> np.ndarray:
    return header.position_start_m + np.arange(header.point_count, dtype=np.float32) * float(header.point_interval_m)


def select_point_indices(header: DatHeader, requested: PointRange) -> tuple[np.ndarray, np.ndarray, str]:
    positions = positions_for_header(header)
    low = max(float(requested.low_m), float(np.min(positions)))
    high = min(float(requested.high_m), float(np.max(positions)))
    warnings: list[str] = []
    if low > high:
        center = (requested.low_m + requested.high_m) / 2.0
        nearest = int(np.argmin(np.abs(positions - center)))
        warnings.append(
            f"请求范围 {requested.text} 超出设备空间 {positions[0]:.2f}-{positions[-1]:.2f} m，已使用最近点 {positions[nearest]:.2f} m"
        )
        return np.array([nearest], dtype=np.int64), positions[[nearest]], "; ".join(warnings)
    if requested.low_m < positions[0] or requested.high_m > positions[-1]:
        warnings.append(f"请求范围 {requested.text} 已裁剪为 {low:.2f}-{high:.2f} m")
    indices = np.flatnonzero((positions >= low - 1e-6) & (positions <= high + 1e-6))
    if indices.size == 0:
        center = (low + high) / 2.0
        nearest = int(np.argmin(np.abs(positions - center)))
        warnings.append(f"范围内无离散点位，已使用最近点 {positions[nearest]:.2f} m")
        indices = np.array([nearest], dtype=np.int64)
    return indices.astype(np.int64), positions[indices], "; ".join(warnings)


def make_bands() -> list[Band]:
    bands = [Band("0Hz", 0.0, 0.0, zero_only=True), Band("1~10Hz", 1.0, 10.0)]
    start = 10.0
    while start < 200.0:
        end = min(start + 15.0, 200.0)
        bands.append(Band(f"{start:g}~{end:g}Hz", start, end))
        start = end
    return bands


def compute_band_energy(
    paths: list[Path],
    header: DatHeader,
    headers: dict[Path, DatHeader],
    point_indices: np.ndarray,
    bands: list[Band],
    window_s: float,
    overlap_percent: float,
) -> tuple[np.ndarray, np.ndarray]:
    sample_rate_hz = max(1, int(header.frequency_hz))
    window_samples = int(round(float(window_s) * sample_rate_hz))
    if window_samples < 2:
        raise ValueError("STFT 时间窗口对应采样点数不足")
    hop_samples = int(round(window_samples * (1.0 - float(overlap_percent) / 100.0)))
    hop_samples = max(1, hop_samples)
    if header.sample_count < window_samples:
        raise ValueError(f"数据长度不足以容纳一个 STFT 窗: {header.sample_count} < {window_samples}")

    freqs = np.fft.rfftfreq(window_samples, d=1.0 / sample_rate_hz).astype(np.float32)
    df = float(freqs[1] - freqs[0]) if freqs.size >= 2 else 1.0
    masks: list[np.ndarray] = []
    for band in bands:
        if band.zero_only:
            masks.append(np.array([0], dtype=np.int64))
            continue
        if band.high_hz > float(freqs[-1]):
            raise ValueError(f"频段 {band.label} 超过 Nyquist 频率 {float(freqs[-1]):.3f} Hz")
        mask = np.flatnonzero((freqs >= band.low_hz) & (freqs <= band.high_hz))
        if mask.size == 0:
            raise ValueError(f"频段 {band.label} 内没有 FFT 频点")
        masks.append(mask.astype(np.int64))

    window_arr = make_stft_window("Hann", window_samples).astype(np.float32, copy=False)
    window_power = float(np.sum(window_arr * window_arr))
    if window_power <= 0.0:
        raise ValueError("窗函数能量为 0，无法计算频段能量")
    scale = compute_silixa_scale(paths, headers) if header.protocol == "Silixa TDMS" else 1.0

    buffer = np.empty((int(point_indices.size), 0), dtype=np.float32)
    buffer_start_global = 0
    loaded_until = 0
    next_start = 0
    columns: list[np.ndarray] = []
    centers: list[float] = []

    for path in paths:
        file_header = headers[path]
        matrix = read_signal_matrix(path, file_header)[point_indices, :].astype(np.float32, copy=False)
        if scale != 1.0:
            matrix = matrix * scale
        buffer = np.concatenate((buffer, matrix), axis=1)
        loaded_until += file_header.sample_count

        while next_start + window_samples <= loaded_until:
            local_start = next_start - buffer_start_global
            local_stop = local_start + window_samples
            block = buffer[:, local_start:local_stop].astype(np.float32, copy=False)
            block = block - np.mean(block, axis=1, keepdims=True)
            spectrum = np.fft.rfft(block * window_arr[np.newaxis, :], axis=1)
            psd = (np.abs(spectrum) ** 2) / (sample_rate_hz * window_power)
            if psd.shape[1] > 1:
                if window_samples % 2 == 0:
                    if psd.shape[1] > 2:
                        psd[:, 1:-1] *= 2.0
                else:
                    psd[:, 1:] *= 2.0

            band_values = np.empty((len(bands), int(point_indices.size)), dtype=np.float32)
            for band_index, (band, mask) in enumerate(zip(bands, masks)):
                energy = np.sum(psd[:, mask], axis=1) * df
                if band.zero_only:
                    energy = psd[:, 0] * df
                band_values[band_index, :] = 10.0 * np.log10(np.maximum(energy, 1e-24))
            columns.append(band_values)
            centers.append((next_start + window_samples / 2.0) / sample_rate_hz)
            next_start += hop_samples

        keep_from_global = next_start
        if keep_from_global > buffer_start_global:
            drop = keep_from_global - buffer_start_global
            buffer = buffer[:, drop:]
            buffer_start_global = keep_from_global

    if not columns:
        raise ValueError("没有生成任何 STFT 时间窗")
    energy = np.stack(columns, axis=2)
    return energy, np.asarray(centers, dtype=np.float32)


def find_font(size: int) -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    candidates = [
        Path(r"C:\Windows\Fonts\msyh.ttc"),
        Path(r"C:\Windows\Fonts\simhei.ttf"),
        Path(r"C:\Windows\Fonts\simsun.ttc"),
        Path(r"C:\Windows\Fonts\arial.ttf"),
    ]
    for candidate in candidates:
        if candidate.exists():
            return ImageFont.truetype(str(candidate), size=size)
    return ImageFont.load_default()


def colorize_phase(values: np.ndarray, vlim: float) -> Image.Image:
    clipped = np.clip(values, -vlim, vlim)
    normalized = (clipped + vlim) / (2.0 * vlim)
    return Image.fromarray(turbo_rgb(normalized), mode="RGB")


def turbo_rgb(normalized: np.ndarray) -> np.ndarray:
    x = np.clip(normalized.astype(np.float32), 0.0, 1.0)
    red = 34.61 + x * (1172.33 + x * (-10793.56 + x * (33300.12 + x * (-38394.49 + x * 14825.05))))
    green = 23.31 + x * (557.33 + x * (1225.33 + x * (-3574.96 + x * (1073.77 + x * 707.56))))
    blue = 27.2 + x * (3211.1 + x * (-15327.97 + x * (27814.0 + x * (-22569.18 + x * 6838.66))))
    rgb = np.stack([red, green, blue], axis=-1)
    return np.uint8(np.clip(rgb, 0, 255))


def colorize_turbo(values: np.ndarray, vmin: float, vmax: float) -> Image.Image:
    if not np.isfinite(vmin) or not np.isfinite(vmax) or vmax <= vmin:
        vmin, vmax = float(np.nanmin(values)), float(np.nanmax(values))
        if vmax <= vmin:
            vmax = vmin + 1.0
    normalized = (values - vmin) / (vmax - vmin)
    normalized = np.nan_to_num(normalized, nan=0.0, posinf=1.0, neginf=0.0)
    return Image.fromarray(turbo_rgb(normalized), mode="RGB")


def draw_centered_text(draw: ImageDraw.ImageDraw, xy: tuple[int, int], text: str, font: ImageFont.ImageFont, fill: str) -> None:
    bbox = draw.textbbox((0, 0), text, font=font)
    width = bbox[2] - bbox[0]
    draw.text((xy[0] - width // 2, xy[1]), text, font=font, fill=fill)


def render_phase_heatmap(
    display: np.ndarray,
    header: DatHeader,
    record: ExperimentRecord,
    device: DeviceConfig,
    paths: list[Path],
    output_path: Path,
    vlim: float,
) -> None:
    image_matrix = display.T
    heatmap = colorize_phase(image_matrix, vlim=vlim)

    canvas_w, canvas_h = 1800, 1120
    left, top, plot_w, plot_h = 145, 165, 1360, 760
    right = left + plot_w
    bottom = top + plot_h
    canvas = Image.new("RGB", (canvas_w, canvas_h), "white")
    draw = ImageDraw.Draw(canvas)

    title_font = find_font(44)
    label_font = find_font(24)
    small_font = find_font(20)
    tick_font = find_font(18)

    draw.text((70, 38), f"{record.title} | {device.title}", font=title_font, fill="#111111")
    subtitle = f"相位热力图 | 全点位空间 {header.position_start_m:.1f}-{header.position_end_m:.1f} m | 色条 ±{vlim:g} | 文件 {len(paths)} 个"
    draw.text((72, 100), subtitle, font=label_font, fill="#444444")

    heatmap = heatmap.resize((plot_w, plot_h), resample=Image.Resampling.BILINEAR)
    canvas.paste(heatmap, (left, top))
    draw.rectangle((left, top, right, bottom), outline="#111111", width=2)

    x_ticks = np.linspace(header.position_start_m, header.position_end_m, 6)
    for tick in x_ticks:
        x = left + int(round((tick - header.position_start_m) / max(header.distance_span_m, header.point_interval_m, 1e-9) * plot_w))
        draw.line((x, bottom, x, bottom + 8), fill="#111111", width=2)
        draw_centered_text(draw, (x, bottom + 12), f"{tick:.0f}", tick_font, "#111111")

    y_ticks = np.linspace(0.0, header.time_span_s, 6)
    for tick in y_ticks:
        y = top + int(round(tick / max(header.time_span_s, 1e-9) * plot_h))
        draw.line((left - 8, y, left, y), fill="#111111", width=2)
        bbox = draw.textbbox((0, 0), f"{tick:.0f}", font=tick_font)
        draw.text((left - 16 - (bbox[2] - bbox[0]), y - 10), f"{tick:.0f}", font=tick_font, fill="#111111")

    draw_centered_text(draw, ((left + right) // 2, bottom + 52), "Distance (m)", label_font, "#111111")
    draw.text((36, (top + bottom) // 2 - 18), "Time (s)", font=label_font, fill="#111111")

    bar_x, bar_y, bar_w, bar_h = right + 72, top, 42, plot_h
    ramp_values = np.linspace(vlim, -vlim, bar_h, dtype=np.float32).reshape(bar_h, 1)
    bar = colorize_phase(ramp_values, vlim=vlim).resize((bar_w, bar_h), resample=Image.Resampling.BILINEAR)
    canvas.paste(bar, (bar_x, bar_y))
    draw.rectangle((bar_x, bar_y, bar_x + bar_w, bar_y + bar_h), outline="#111111", width=1)
    for value in (vlim, 0.0, -vlim):
        y = bar_y + int(round((vlim - value) / (2.0 * vlim) * bar_h))
        draw.line((bar_x + bar_w, y, bar_x + bar_w + 8, y), fill="#111111", width=2)
        draw.text((bar_x + bar_w + 14, y - 12), f"{value:.2f}", font=tick_font, fill="#111111")
    draw.text((bar_x - 8, bar_y - 36), "Phase", font=label_font, fill="#111111")

    draw.text((70, canvas_h - 92), f"{paths[0].name}  ->  {paths[-1].name}", font=small_font, fill="#333333")
    draw.text((70, canvas_h - 54), f"数据源: {paths[0].parent}", font=small_font, fill="#555555")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output_path)


def render_band_energy_heatmap(
    energy_db: np.ndarray,
    positions: np.ndarray,
    centers_s: np.ndarray,
    band: Band,
    record: ExperimentRecord,
    device: DeviceConfig,
    requested_range: PointRange,
    output_path: Path,
    warnings: str,
    fixed_vmin: float | None = None,
    fixed_vmax: float | None = None,
    robust_levels: bool = False,
) -> None:
    values = np.asarray(energy_db, dtype=np.float32)
    if fixed_vmin is not None and fixed_vmax is not None:
        vmin = float(fixed_vmin)
        vmax = float(fixed_vmax)
        level_note = f"固定色阶 = {vmin:.1f} ~ {vmax:.1f} dB"
    elif robust_levels:
        finite = values[np.isfinite(values)]
        if finite.size:
            valid = finite[finite > -200.0]
            if valid.size < max(16, int(finite.size * 0.01)):
                valid = finite
            vmin = float(np.percentile(valid, 2.0))
            vmax = float(np.percentile(valid, 99.5))
        else:
            vmin, vmax = -120.0, 0.0
        level_note = f"稳健自适应色阶: P2-P99.5 = {vmin:.1f} ~ {vmax:.1f} dB"
    else:
        finite = values[np.isfinite(values)]
        if finite.size:
            vmin = float(np.percentile(finite, 1.0))
            vmax = float(np.percentile(finite, 99.5))
        else:
            vmin, vmax = -120.0, 0.0
        level_note = f"色阶: P1-P99.5 = {vmin:.1f} ~ {vmax:.1f} dB"
    if vmax <= vmin:
        vmax = vmin + 1.0

    heatmap = colorize_turbo(values, vmin, vmax)
    canvas_w, canvas_h = 1800, 1120
    left, top, plot_w, plot_h = 160, 170, 1345, 750
    right = left + plot_w
    bottom = top + plot_h
    canvas = Image.new("RGB", (canvas_w, canvas_h), "white")
    draw = ImageDraw.Draw(canvas)

    title_font = find_font(42)
    label_font = find_font(24)
    small_font = find_font(20)
    tick_font = find_font(18)

    draw.text((70, 38), f"{record.title} | {device.title}", font=title_font, fill="#111111")
    actual_range = f"{float(positions[0]):.1f}-{float(positions[-1]):.1f} m"
    subtitle = f"频谱能量图 {band.label} | STFT 2s, overlap 50% | 请求 {requested_range.text} | 实际 {actual_range}"
    draw.text((72, 100), subtitle, font=label_font, fill="#444444")

    heatmap = heatmap.resize((plot_w, plot_h), resample=Image.Resampling.BILINEAR)
    canvas.paste(heatmap, (left, top))
    draw.rectangle((left, top, right, bottom), outline="#111111", width=2)

    x_min = float(centers_s[0])
    x_max = float(centers_s[-1]) if centers_s.size > 1 else x_min + 1.0
    for tick in np.linspace(x_min, x_max, 6):
        x = left + int(round((tick - x_min) / max(x_max - x_min, 1e-9) * plot_w))
        draw.line((x, bottom, x, bottom + 8), fill="#111111", width=2)
        draw_centered_text(draw, (x, bottom + 12), f"{tick:.0f}", tick_font, "#111111")

    y_min = float(positions[0])
    y_max = float(positions[-1]) if positions.size > 1 else y_min + 1.0
    for tick in np.linspace(y_min, y_max, 6):
        y = top + int(round((tick - y_min) / max(y_max - y_min, 1e-9) * plot_h))
        draw.line((left - 8, y, left, y), fill="#111111", width=2)
        bbox = draw.textbbox((0, 0), f"{tick:.0f}", font=tick_font)
        draw.text((left - 16 - (bbox[2] - bbox[0]), y - 10), f"{tick:.0f}", font=tick_font, fill="#111111")

    draw_centered_text(draw, ((left + right) // 2, bottom + 52), "Time (s)", label_font, "#111111")
    draw.text((36, (top + bottom) // 2 - 18), "Distance (m)", font=label_font, fill="#111111")

    bar_x, bar_y, bar_w, bar_h = right + 72, top, 42, plot_h
    ramp_values = np.linspace(vmax, vmin, bar_h, dtype=np.float32).reshape(bar_h, 1)
    bar = colorize_turbo(ramp_values, vmin, vmax).resize((bar_w, bar_h), resample=Image.Resampling.BILINEAR)
    canvas.paste(bar, (bar_x, bar_y))
    draw.rectangle((bar_x, bar_y, bar_x + bar_w, bar_y + bar_h), outline="#111111", width=1)
    for value in (vmax, (vmin + vmax) / 2.0, vmin):
        y = bar_y + int(round((vmax - value) / max(vmax - vmin, 1e-9) * bar_h))
        draw.line((bar_x + bar_w, y, bar_x + bar_w + 8, y), fill="#111111", width=2)
        draw.text((bar_x + bar_w + 14, y - 12), f"{value:.1f}", font=tick_font, fill="#111111")
    draw.text((bar_x - 66, bar_y - 36), "dB re rad^2", font=label_font, fill="#111111")

    note = warnings or "点位范围与设备空间匹配"
    draw.text((70, canvas_h - 82), note, font=small_font, fill="#555555")
    draw.text((70, canvas_h - 48), level_note, font=small_font, fill="#555555")

    output_path.parent.mkdir(parents=True, exist_ok=True)
    canvas.save(output_path)


def safe_filename(text: str) -> str:
    return re.sub(r'[\\/:*?"<>|\s]+', "_", text).strip("_")


def write_summary_csv(path: Path, summaries: Iterable[DeviceRecordSummary]) -> None:
    rows = [asdict(item) for item in summaries]
    path.parent.mkdir(parents=True, exist_ok=True)
    fieldnames = list(DeviceRecordSummary.__annotations__.keys())
    with path.open("w", newline="", encoding="utf-8-sig") as f:
        writer = csv.DictWriter(f, fieldnames=fieldnames)
        writer.writeheader()
        writer.writerows(rows)


def output_folder_for(output_dir: Path, record: ExperimentRecord, device: DeviceConfig) -> Path:
    return output_dir / safe_filename(record.group) / record.folder_name / safe_filename(device.title)


def process_record_device(
    record: ExperimentRecord,
    device: DeviceConfig,
    files: list[DatFile],
    output_dir: Path,
    args: argparse.Namespace,
    bands: list[Band],
) -> DeviceRecordSummary:
    requested_range = record.device_ranges[device.title]
    matched = files_for_record(files, record)
    device_out = output_folder_for(output_dir, record, device)
    summary = DeviceRecordSummary(
        group=record.group,
        seq=record.seq,
        device=device.title,
        start=record.start.isoformat(sep=" "),
        end=record.end.isoformat(sep=" "),
        operation=record.operation,
        flow_lh=record.flow_lh,
        requested_range_m=requested_range.text,
        file_count=len(matched),
        first_file=matched[0].path.name if matched else "",
        last_file=matched[-1].path.name if matched else "",
        output_folder=str(device_out),
    )
    if not matched:
        summary.message = "时间范围内没有匹配到数据文件"
        return summary
    if args.match_only:
        summary.status = "matched"
        return summary

    paths = [item.path for item in matched]
    try:
        header, headers = validate_headers(paths)
        indices, positions, range_warning = select_point_indices(header, requested_range)
        summary.protocol = header.protocol
        summary.frequency_hz = header.frequency_hz
        summary.point_count = header.point_count
        summary.position_start_m = round(header.position_start_m, 3)
        summary.position_end_m = round(header.position_end_m, 3)
        summary.sample_count = header.sample_count
        summary.duration_s = round(header.time_span_s, 3)
        summary.selected_point_count = int(indices.size)
        summary.actual_range_m = f"{float(positions[0]):.3f}-{float(positions[-1]):.3f}"
        summary.warnings = range_warning

        if not args.skip_phase:
            phase_output = device_out / "phase" / f"{record.folder_name}_{safe_filename(device.title)}_phase_heatmap.png"
            display, phase_header = load_phase_display_matrix(paths, max_time_bins=args.max_time_bins)
            render_phase_heatmap(display, phase_header, record, device, paths, phase_output, vlim=args.phase_vlim)
            summary.phase_image = str(phase_output)

        if not args.skip_bands:
            energy_by_band, centers_s = compute_band_energy(
                paths,
                header,
                headers,
                indices,
                bands,
                window_s=args.stft_window_s,
                overlap_percent=args.stft_overlap,
            )
            band_dir = device_out / "band_energy"
            for band_index, band in enumerate(bands):
                band_output = band_dir / f"{record.folder_name}_{safe_filename(device.title)}_{band.filename_label}_energy.png"
                render_band_energy_heatmap(
                    energy_by_band[band_index],
                    positions,
                    centers_s,
                    band,
                    record,
                    device,
                    requested_range,
                    band_output,
                    warnings=range_warning,
                    fixed_vmin=args.band_vmin,
                    fixed_vmax=args.band_vmax,
                    robust_levels=args.band_robust_levels,
                )
            summary.band_image_count = len(bands)

        summary.status = "ok"
    except Exception as exc:
        summary.status = "failed"
        summary.message = str(exc)
    return summary


def process(args: argparse.Namespace) -> int:
    records = read_records(args.workbook)
    output_dir = args.output_dir
    output_dir.mkdir(parents=True, exist_ok=True)
    bands = make_bands()

    selected_devices = [
        config for config in DEVICE_CONFIGS
        if not args.device or config.title.lower() == args.device.lower() or config.key.lower() == args.device.lower()
    ]
    if not selected_devices:
        raise SystemExit(f"没有找到设备: {args.device}")

    files_by_device: dict[str, list[DatFile]] = {}
    for config in selected_devices:
        data_root = args.data_root / config.data_subdir
        files = index_data_files(data_root)
        files_by_device[config.title] = files
        if files:
            print(
                f"{config.title}: {len(files)} files, "
                f"{files[0].timestamp:%Y-%m-%d %H:%M:%S} ~ {files[-1].timestamp:%Y-%m-%d %H:%M:%S}"
            )
        else:
            print(f"{config.title}: no data files under {data_root}")

    print(f"records: {len(records)}")
    print("bands: " + ", ".join(band.label for band in bands))
    summaries: list[DeviceRecordSummary] = []
    total = len(records) * len(selected_devices)
    done = 0
    for record in records:
        for config in selected_devices:
            done += 1
            print(f"[{done}/{total}] {record.group} {record.seq} {config.title} {record.start:%H:%M:%S}-{record.end:%H:%M:%S}")
            summary = process_record_device(record, config, files_by_device[config.title], output_dir, args, bands)
            summaries.append(summary)
            print(f"  status={summary.status}, files={summary.file_count}, bands={summary.band_image_count}, {summary.message}")

    if not args.no_summary:
        csv_path = output_dir / "well_flow_three_device_summary.csv"
        json_path = output_dir / "well_flow_three_device_summary.json"
        write_summary_csv(csv_path, summaries)
        json_path.write_text(json.dumps([asdict(item) for item in summaries], ensure_ascii=False, indent=2), encoding="utf-8")
        print(f"SUMMARY_CSV {csv_path}")
        print(f"SUMMARY_JSON {json_path}")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Batch well-flow phase and band-energy heatmaps for three TDAS devices.")
    parser.add_argument("--workbook", type=Path, default=DEFAULT_WORKBOOK)
    parser.add_argument("--data-root", type=Path, default=DEFAULT_DATA_ROOT)
    parser.add_argument("--output-dir", type=Path, default=DEFAULT_OUTPUT_DIR)
    parser.add_argument("--phase-vlim", type=float, default=DEFAULT_PHASE_VLIM)
    parser.add_argument("--max-time-bins", type=int, default=DEFAULT_MAX_TIME_BINS)
    parser.add_argument("--stft-window-s", type=float, default=DEFAULT_STFT_WINDOW_S)
    parser.add_argument("--stft-overlap", type=float, default=DEFAULT_STFT_OVERLAP_PERCENT)
    parser.add_argument("--match-only", action="store_true")
    parser.add_argument("--skip-phase", action="store_true")
    parser.add_argument("--skip-bands", action="store_true")
    parser.add_argument("--no-summary", action="store_true")
    parser.add_argument("--device", type=str, default="", help="只处理指定设备标题或key，如 OV、silixia、self_2u。")
    parser.add_argument("--band-vmin", type=float, default=None, help="频段能量图固定色条下限。")
    parser.add_argument("--band-vmax", type=float, default=None, help="频段能量图固定色条上限。")
    parser.add_argument("--band-robust-levels", action="store_true", help="频段能量图使用稳健自适应色阶。")
    args = parser.parse_args()
    return process(args)


if __name__ == "__main__":
    raise SystemExit(main())
