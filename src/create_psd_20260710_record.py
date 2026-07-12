from __future__ import annotations

import sys
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

PROJECT_ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(PROJECT_ROOT / "src"))

from batch_well_flow_three_device_analysis import index_data_files  # noqa: E402


OUTPUT_DIR = PROJECT_ROOT / "井中流动分析" / "PSD分析_20260710"
OUTPUT_XLSX = OUTPUT_DIR / "中国石油大学PSD分析法_实验记录与数据覆盖_已确认.xlsx"
DATA_ROOT = Path(r"E:\北京试验")

DEVICE_FOLDERS = {
    "silixia": "silixia",
    "OV": "ov",
    "自研_2U": "自研_2U",
}

PSD_BANDS = "0~2Hz; 2~10Hz; 10~50Hz; 50~200Hz"
SILIXIA_RANGE = "1000~1130m"
OV_RANGE = "0~396m"
SELF_2U_RANGE = "0~370m"


@dataclass(frozen=True)
class ExperimentGroup:
    date: str
    test_name: str
    group_id: str
    start: str
    end: str
    noise_start: str = ""
    noise_end: str = ""
    pump_start: str = ""
    pump_end: str = ""
    stable_start: str = ""
    stable_end: str = ""
    stop_start: str = ""
    stop_end: str = ""
    liquid_flow_m3h: str = ""
    liquid_flow_lh: str = ""
    gas_or_ratio: str = ""
    pulse_ns: str = "50"
    sample_rate_khz: str = "4"
    spatial_spacing_m: str = "5"
    silixia_range: str = SILIXIA_RANGE
    ov_range: str = OV_RANGE
    self_2u_range: str = SELF_2U_RANGE
    psd_bands: str = PSD_BANDS
    stft_window_s: str = "2"
    stft_overlap_percent: str = "50"
    source_image: str = ""
    qc_status: str = "已确认"
    note: str = ""

    @property
    def phase_summary(self) -> str:
        phases: list[str] = []
        for label, start, end in (
            ("底噪", self.noise_start, self.noise_end),
            ("调泵", self.pump_start, self.pump_end),
            ("稳定", self.stable_start, self.stable_end),
            ("停止", self.stop_start, self.stop_end),
        ):
            if start and end:
                phases.append(f"{label} {start}-{end}")
        return "; ".join(phases)


def hhmm_to_dt(date_text: str, time_text: str) -> datetime | None:
    clean = time_text.strip()
    if not clean:
        return None
    for fmt in ("%Y-%m-%d %H:%M", "%Y-%m-%d %H:%M:%S"):
        try:
            return datetime.strptime(f"{date_text} {clean}", fmt)
        except ValueError:
            continue
    return None


def make_records() -> list[ExperimentGroup]:
    date = "2026-07-10"
    img1 = "73c46a1f7ec71e4e214a35104c3ee2f6(1).jpg"
    img2 = "f2e40b44a53e2b20737f5e70a8ec2209(1).jpg"
    return [
        ExperimentGroup(
            date, "大流量测试", "W1_2.5m3h", "10:56", "11:07",
            noise_start="10:56", noise_end="11:01",
            pump_start="11:01", pump_end="11:02",
            stable_start="11:02", stable_end="11:07",
            liquid_flow_m3h="2.5", liquid_flow_lh="2500", source_image=img1,
            note="第一组按完整实验组绘图，含底噪、调泵和稳定三个阶段。",
        ),
        ExperimentGroup(
            date, "大流量测试", "W2_2.0m3h", "11:07", "11:14",
            pump_start="11:07", pump_end="11:09",
            stable_start="11:09", stable_end="11:14",
            liquid_flow_m3h="2.0", liquid_flow_lh="2000", source_image=img1,
        ),
        ExperimentGroup(
            date, "大流量测试", "W3_1.5m3h", "11:14", "11:21",
            pump_start="11:14", pump_end="11:16",
            stable_start="11:16", stable_end="11:21",
            liquid_flow_m3h="1.5", liquid_flow_lh="1500", source_image=img1,
        ),
        ExperimentGroup(
            date, "大流量测试", "W4_1.0m3h", "11:21", "11:28",
            pump_start="11:21", pump_end="11:23",
            stable_start="11:23", stable_end="11:28",
            liquid_flow_m3h="1.0", liquid_flow_lh="1000", source_image=img1,
        ),
        ExperimentGroup(
            date, "大流量测试", "W5_0.5m3h", "11:28", "11:38",
            pump_start="11:28", pump_end="11:30",
            stable_start="11:30", stable_end="11:35",
            stop_start="11:35", stop_end="11:38",
            liquid_flow_m3h="0.5", liquid_flow_lh="500", source_image=img1,
            note="最后一组按完整实验组绘图，包含调泵、稳定和停止阶段。",
        ),
        ExperimentGroup(
            date, "气水两相实验", "G1_纯水100%", "14:55", "15:07",
            noise_start="14:55", noise_end="15:00",
            pump_start="15:00", pump_end="15:02",
            stable_start="15:02", stable_end="15:07",
            liquid_flow_m3h="1.0", liquid_flow_lh="1000", gas_or_ratio="纯水100%",
            source_image=img2,
            note="第一组按完整实验组绘图，含底噪、调泵和稳定三个阶段。",
        ),
        ExperimentGroup(
            date, "气水两相实验", "G2_气水比75%", "15:15", "15:25",
            pump_start="15:15", pump_end="15:20",
            stable_start="15:20", stable_end="15:25",
            liquid_flow_m3h="1.0", liquid_flow_lh="1000", gas_or_ratio="气水比75%",
            source_image=img2,
            note="调泵起始时间15:15及15:15-15:20调泵区间已由用户确认。",
        ),
        ExperimentGroup(
            date, "气水两相实验", "G3_气水比50%", "15:25", "15:33",
            pump_start="15:25", pump_end="15:28",
            stable_start="15:28", stable_end="15:33",
            liquid_flow_m3h="1.0", liquid_flow_lh="1000", gas_or_ratio="气水比50%",
            source_image=img2,
        ),
        ExperimentGroup(
            date, "气水两相实验", "G4_气水比25%", "15:33", "15:41",
            pump_start="15:33", pump_end="15:36",
            stable_start="15:36", stable_end="15:41",
            liquid_flow_m3h="1.0", liquid_flow_lh="1000", gas_or_ratio="气水比25%",
            source_image=img2,
        ),
        ExperimentGroup(
            date, "气水两相实验", "G5_单相气3m3h", "15:41", "15:56",
            pump_start="15:41", pump_end="15:51",
            stable_start="15:51", stable_end="15:56",
            gas_or_ratio="3m3/h单相气", source_image=img2,
            note="15:41-15:51调泵区间已确认；停止阶段无明确结束时间，按此前确认不纳入。",
        ),
    ]


HEADERS = [
    "日期", "实验类型", "实验编号", "实验开始时间", "实验结束时间", "阶段说明",
    "底噪开始", "底噪结束", "调泵开始", "调泵结束", "稳定开始", "稳定结束", "停止开始", "停止结束",
    "液体流量(m3/h)", "液体流量(L/h)", "气量/气水比说明",
    "silixia频谱点位范围(m)", "OV频谱点位范围(m)", "自研_2U频谱点位范围(m)",
    "脉宽(ns)", "采样率(kHz)", "空间间距(m)", "PSD频带", "STFT窗口(s)", "STFT重叠度(%)",
    "来源图片", "核对状态", "备注",
]


def record_to_row(record: ExperimentGroup) -> list[Any]:
    return [
        record.date, record.test_name, record.group_id, record.start, record.end, record.phase_summary,
        record.noise_start, record.noise_end, record.pump_start, record.pump_end,
        record.stable_start, record.stable_end, record.stop_start, record.stop_end,
        record.liquid_flow_m3h, record.liquid_flow_lh, record.gas_or_ratio,
        record.silixia_range, record.ov_range, record.self_2u_range,
        record.pulse_ns, record.sample_rate_khz, record.spatial_spacing_m,
        record.psd_bands, record.stft_window_s, record.stft_overlap_percent,
        record.source_image, record.qc_status, record.note,
    ]


def scan_coverage(records: list[ExperimentGroup]) -> tuple[list[list[Any]], dict[str, tuple[int, str, str]]]:
    indexed: dict[str, list[Any]] = {}
    device_summary: dict[str, tuple[int, str, str]] = {}
    for device, folder in DEVICE_FOLDERS.items():
        data_dir = DATA_ROOT / folder
        files = index_data_files(data_dir) if data_dir.exists() else []
        indexed[device] = files
        if files:
            device_summary[device] = (
                len(files), files[0].timestamp.isoformat(sep=" "), files[-1].timestamp.isoformat(sep=" ")
            )
        else:
            device_summary[device] = (0, "", "")

    coverage_rows: list[list[Any]] = []
    for record in records:
        start_dt = hhmm_to_dt(record.date, record.start)
        end_dt = hhmm_to_dt(record.date, record.end)
        for device, files in indexed.items():
            matched = [item for item in files if start_dt and end_dt and start_dt <= item.timestamp <= end_dt]
            if matched:
                status = "有匹配数据"
                message = "按完整实验组时间范围匹配；正式绘图会读取头信息复核采样率和空间范围。"
                first = matched[0].timestamp.isoformat(sep=" ")
                last = matched[-1].timestamp.isoformat(sep=" ")
            else:
                status = "未找到匹配数据"
                message = "该完整实验组时间范围内未找到文件名时间戳匹配的数据。"
                first = ""
                last = ""
            coverage_rows.append([
                record.date, record.test_name, record.group_id, record.phase_summary, device,
                record.start, record.end, len(matched), first, last, status, message,
            ])
    return coverage_rows, device_summary


def style_sheet(ws) -> None:
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)
    pending_fill = PatternFill("solid", fgColor="FFF2CC")
    for cell in ws[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if any("待" in str(cell.value or "") for cell in row):
            for cell in row:
                cell.fill = pending_fill
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    ws.row_dimensions[1].height = 34
    for column_cells in ws.columns:
        max_len = max(len(str(cell.value or "")) for cell in column_cells)
        ws.column_dimensions[get_column_letter(column_cells[0].column)].width = min(max(max_len + 2, 10), 36)


def build_workbook() -> Path:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    records = make_records()
    coverage_rows, device_summary = scan_coverage(records)

    workbook = Workbook()
    ws = workbook.active
    ws.title = "实验记录_已确认"
    ws.append(HEADERS)
    for record in records:
        ws.append(record_to_row(record))
    style_sheet(ws)

    ws_cov = workbook.create_sheet("数据覆盖检查")
    ws_cov.append([
        "日期", "实验类型", "实验编号", "组内阶段", "设备", "请求开始", "请求结束",
        "匹配文件数", "匹配首文件时间", "匹配末文件时间", "覆盖状态", "说明",
    ])
    for row in coverage_rows:
        ws_cov.append(row)
    style_sheet(ws_cov)

    ws_dev = workbook.create_sheet("设备目录概览")
    ws_dev.append(["设备", "目录", "文件数", "最早文件时间", "最晚文件时间"])
    for device, folder in DEVICE_FOLDERS.items():
        count, first, last = device_summary[device]
        ws_dev.append([device, str(DATA_ROOT / folder), count, first, last])
    style_sheet(ws_dev)

    ws_qc = workbook.create_sheet("人工核对清单")
    ws_qc.append(["序号", "核对项", "状态", "说明"])
    checks = [
        ("1", "实验记录粒度", "已确认", "一行一个完整实验组；阶段仅作为组内时间标记，不拆分绘图。"),
        ("2", "首组阶段", "已确认", "每类实验首组包含底噪、调泵和稳定。"),
        ("3", "中间组阶段", "已确认", "中间实验组只包含调泵和稳定。"),
        ("4", "末组停止阶段", "已确认", "大流量W5包含停止；气水两相G5因停止结束时间不清楚，按此前确认不纳入。"),
        ("5", "三设备频谱点位", "已确认", "silixia 1000~1130m；OV 0~396m；自研_2U 0~370m。"),
        ("6", "PSD设置", "已确认", "0~2、2~10、10~50、50~200Hz；Hann窗2s；重叠50%。"),
        ("7", "相位热力图设置", "已确认", "完整实验组、全空间点位、固定色条-0.15~0.15，并标记阶段边界。"),
        ("8", "原始数据覆盖", "待复核", "表内先按文件名时间戳检查；正式绘图继续读取文件头和实际数据。"),
    ]
    for row in checks:
        ws_qc.append(row)
    style_sheet(ws_qc)

    workbook.save(OUTPUT_XLSX)
    return OUTPUT_XLSX


if __name__ == "__main__":
    print(build_workbook())
